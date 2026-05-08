from __future__ import annotations

import base64
import re
import shutil
import subprocess
from io import BytesIO
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any
from urllib.parse import quote

import httpx
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches as DocxInches, Pt, RGBColor
from playwright.sync_api import Error as PlaywrightError
from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.dml.color import RGBColor as PptxRGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt as PptxPt
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfgen import canvas as pdf_canvas
from reportlab.platypus import ListFlowable, ListItem, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .config import Settings


@dataclass
class Block:
    kind: str
    text: str
    level: int = 0
    rows: list[list[str]] | None = None


@dataclass
class Slide:
    title: str
    kicker: str = ""
    body: list[str] | None = None
    bullets: list[str] | None = None
    rows: list[list[str]] | None = None
    layout: str = "content"
    visual: str | None = None


@dataclass(frozen=True)
class DeckVisualDirection:
    domain: str
    mood: str
    dark_bg: str
    dark_divider: str
    light_bg: str
    card_bg: str
    ink: str
    muted: str
    accent: str
    accent_2: str
    line: str
    dark_muted: str
    dark_line: str
    cover_label: str
    focus_label: str
    image_style: str


class DocumentService:
    def __init__(self, settings: Settings) -> None:
        self.settings = settings
        self.output_root = Path("generated_documents")

    def generate(
        self,
        session_id: str,
        job_id: str,
        user_request: str,
        draft: dict[str, Any],
        plan: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        source_text = self._source_text(user_request, draft)
        title = self._title_from_text(user_request, source_text)
        plan_kind = str((plan or {}).get("document_kind") or "").strip().lower()
        is_slide = plan_kind == "slides" or (not plan_kind and self._is_slide_request(user_request))
        if is_slide:
            return self._generate_slides(session_id, job_id, user_request, title, source_text, draft, plan=plan)
        blocks = self._parse_blocks(source_text)
        folder = self.output_root / session_id / job_id
        folder.mkdir(parents=True, exist_ok=True)

        file_stem = self._filename_stem(user_request, title)
        docx_name = self._safe_filename(file_stem, ".docx")
        pdf_name = self._safe_filename(file_stem, ".pdf")
        docx_path = folder / docx_name
        pdf_path = folder / pdf_name

        self._build_docx(docx_path, title, blocks, user_request)
        pdf_provider = self._build_pdf(pdf_path, title, blocks, user_request, session_id)

        summary = "Documento organizado em DOCX e PDF gerado com sucesso."
        raw = (
            f"Arquivos gerados com sucesso: `{docx_name}` e `{pdf_name}`.\n\n"
            "Use os links do chat para abrir ou baixar o DOCX e o PDF."
        )
        return {
            "provider": draft.get("provider", "kimi-documento"),
            "model": f"{draft.get('model', 'kimi-documento')} + {pdf_provider}",
            "summary": summary,
            "raw": raw,
            "document_title": title,
            "files": [
                {
                    "name": docx_name,
                    "path": str(docx_path).replace("\\", "/"),
                    "mime_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    "download_url": f"/api/artefatos/{job_id}/{docx_name}",
                },
                {
                    "name": pdf_name,
                    "path": str(pdf_path).replace("\\", "/"),
                    "mime_type": "application/pdf",
                    "download_url": f"/api/artefatos/{job_id}/{pdf_name}",
                },
            ],
        }

    def _generate_slides(
        self,
        session_id: str,
        job_id: str,
        user_request: str,
        title: str,
        source_text: str,
        draft: dict[str, Any],
        plan: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        folder = self.output_root / session_id / job_id
        folder.mkdir(parents=True, exist_ok=True)

        file_stem = self._filename_stem(user_request, title)
        pptx_name = self._safe_filename(file_stem, ".pptx")
        html_name = self._safe_filename(file_stem, ".slides.html")
        markdown_name = self._safe_filename(file_stem, ".slides.md")
        pdf_name = self._safe_filename(file_stem, ".pdf")
        pptx_path = folder / pptx_name
        html_path = folder / html_name
        markdown_path = folder / markdown_name
        pdf_path = folder / pdf_name

        normalized_source = self._normalize_slide_source(source_text)
        if self._needs_slide_fallback(normalized_source, user_request):
            normalized_source = self._fallback_slide_deck_text(user_request, title)
        raw_slides = self._split_raw_slides(normalized_source) if self._looks_like_marp_deck(normalized_source.strip()) else self._split_into_slides(normalized_source.strip())
        if len(raw_slides) < 6:
            normalized_source = self._fallback_slide_deck_text(user_request, title)
            raw_slides = self._split_into_slides(normalized_source)
        slide_visuals = self._generate_slide_visuals(title, user_request, raw_slides, folder)
        slides = self._build_slide_models(raw_slides, title, slide_visuals)
        pptx_provider = self._build_presentation_pptx(pptx_path, title, slides, folder, user_request)
        html_output = self._build_presentation_html(title, slides, user_request)
        html_output = self._inline_slide_assets(html_output, folder)
        html_path.write_text(html_output, encoding="utf-8")
        markdown_path.write_text(self._build_marp_markdown(title, normalized_source, slide_visuals), encoding="utf-8")

        pdf_provider = ""
        should_render_pdf = True
        if should_render_pdf:
            try:
                pdf_provider = self._build_presentation_pdf(markdown_path, html_path, pdf_path, slides)
            except Exception:
                try:
                    pdf_provider = self._build_presentation_pdf_from_pptx(pptx_path, pdf_path, slides)
                except Exception:
                    pdf_provider = ""

        files: list[dict[str, Any]] = []
        if pptx_path.exists():
            files.append(
                {
                    "name": pptx_name,
                    "path": str(pptx_path).replace("\\", "/"),
                    "mime_type": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                    "download_url": f"/api/artefatos/{job_id}/{pptx_name}",
                }
            )
        if html_path.exists():
            files.append(
                {
                    "name": html_name,
                    "path": str(html_path).replace("\\", "/"),
                    "mime_type": "text/html",
                    "download_url": f"/api/artefatos/{job_id}/{html_name}",
                    "language": "html",
                }
            )
        if pdf_path.exists():
            files.append(
                {
                    "name": pdf_name,
                    "path": str(pdf_path).replace("\\", "/"),
                    "mime_type": "application/pdf",
                    "download_url": f"/api/artefatos/{job_id}/{pdf_name}",
                }
            )
        if markdown_path.exists():
            files.append(
                {
                    "name": markdown_name,
                    "path": str(markdown_path).replace("\\", "/"),
                    "mime_type": "text/markdown",
                    "download_url": f"/api/artefatos/{job_id}/{markdown_name}",
                    "language": "markdown",
                }
            )

        summary = "Apresentacao profissional gerada em PPTX."
        if slide_visuals:
            summary = "Apresentacao profissional gerada em PPTX com composicao visual e imagens IA."
        if html_path.exists():
            summary += " Preview HTML incluido para abrir no painel lateral."
        if pdf_path.exists():
            summary += " PDF de apresentacao incluido."
        raw = f"Apresentacao gerada com sucesso: `{pptx_name}`."
        if html_path.exists():
            raw += f" Arquivo de preview: `{html_name}`."
        if pdf_path.exists():
            raw += f" PDF: `{pdf_name}`."
        if slide_visuals:
            raw += f" Imagens IA aplicadas em {len(slide_visuals)} slide(s)."
        return {
            "provider": draft.get("provider", "kimi-slides"),
            "model": " + ".join(
                part
                for part in [draft.get("model", "kimi-slides"), pptx_provider, pdf_provider]
                if part
            ),
            "summary": summary,
            "raw": raw,
            "document_title": title,
            "files": files,
            "preview_url": f"/api/artefatos/{job_id}/{html_name}" if html_path.exists() else "",
            "slide_deck": True,
            "tools_used": [
                "python-pptx",
                "slide-html-preview",
                *(["pollinations-image"] if slide_visuals else []),
                *(["presentation-pdf-render"] if pdf_path.exists() else []),
            ],
            "execution_contract": (plan or {}).get("response_contract"),
        }

    def _source_text(self, user_request: str, draft: dict[str, Any]) -> str:
        for key in ("raw", "summary"):
            value = str(draft.get(key) or "").strip()
            cleaned = self._normalize_document_text(self._strip_system_notices(value), user_request)
            if self._looks_like_document_content(cleaned):
                return cleaned
        return self._fallback_document_text(user_request)

    def _is_slide_request(self, user_request: str, text: str = "") -> bool:
        lowered = user_request.lower()
        if any(marker in lowered for marker in ["docx", "word", "abnt", "documento", "relatorio", "relatório", "pdf"]) and not any(
            marker in lowered for marker in ["slide", "slides", "deck", "ppt", "pptx", "powerpoint", "apresentacao", "apresentação"]
        ):
            return False
        markers = [
            "slide",
            "slides",
            "apresentacao",
            "apresentação",
            "deck",
            "powerpoint",
            "ppt",
            "palestra",
        ]
        return any(marker in lowered for marker in markers)

    def _title_from_text(self, user_request: str, text: str) -> str:
        extracted = self._extract_topic(user_request)
        if extracted:
            return extracted[:80].rstrip(" .:-")
        for line in text.splitlines():
            cleaned = re.sub(r"^#+\s*", "", line).strip()
            if len(cleaned) >= 6:
                return cleaned[:80]
        base = " ".join(user_request.split()).strip()
        return (base[:80] or "Documento Kimi AI").rstrip(" .:-")

    def _build_marp_markdown(self, title: str, source_text: str, slide_visuals: dict[int, str] | None = None) -> str:
        cleaned = self._normalize_slide_source(source_text).strip()
        slides = self._split_raw_slides(cleaned) if self._looks_like_marp_deck(cleaned) else self._split_into_slides(cleaned)
        slides = self._polish_slide_deck(slides, title, slide_visuals or {})
        frontmatter = (
            "---\n"
            "marp: true\n"
            "theme: gaia\n"
            "paginate: true\n"
            "size: 16:9\n"
            "headingDivider: 2\n"
            "footer: 'Kemy AI'\n"
            "style: |\n"
            "  section {\n"
            "    font-family: 'Aptos', 'Segoe UI', sans-serif;\n"
            "    background: linear-gradient(180deg, #f8fbff 0%, #eef3f9 100%);\n"
            "    color: #10243d;\n"
            "    padding: 54px 64px 58px;\n"
            "  }\n"
            "  section strong {\n"
            "    color: #0d315f;\n"
            "  }\n"
            "  section::after {\n"
            "    font-size: 0.72rem;\n"
            "    color: #68809f;\n"
            "  }\n"
            "  h1 {\n"
            "    color: #0b2e59;\n"
            "    font-size: 2.2rem;\n"
            "    line-height: 1.02;\n"
            "    letter-spacing: -0.03em;\n"
            "    margin: 0 0 0.32em;\n"
            "  }\n"
            "  h2 {\n"
            "    color: #0f437a;\n"
            "    font-size: 1.32rem;\n"
            "    line-height: 1.1;\n"
            "    letter-spacing: -0.02em;\n"
            "    margin: 0 0 0.5em;\n"
            "  }\n"
            "  h3 {\n"
            "    color: #1d5d96;\n"
            "    font-size: 1.0rem;\n"
            "    margin: 0 0 0.5em;\n"
            "  }\n"
            "  p {\n"
            "    font-size: 1rem;\n"
            "    line-height: 1.42;\n"
            "    margin: 0 0 0.7em;\n"
            "  }\n"
            "  ul, ol {\n"
            "    margin: 0.2em 0 0;\n"
            "    padding-left: 1.1em;\n"
            "  }\n"
            "  li {\n"
            "    font-size: 0.96rem;\n"
            "    line-height: 1.38;\n"
            "    margin: 0 0 0.34em;\n"
            "  }\n"
            "  strong {\n"
            "    color: #0b3264;\n"
            "    font-weight: 750;\n"
            "  }\n"
            "  blockquote {\n"
            "    margin: 0.8em 0;\n"
            "    padding: 0.7em 0.9em;\n"
            "    border-left: 5px solid #2f6fb1;\n"
            "    background: rgba(255,255,255,0.62);\n"
            "    border-radius: 0 12px 12px 0;\n"
            "  }\n"
            "  table {\n"
            "    width: 100%;\n"
            "    border-collapse: collapse;\n"
            "    margin-top: 0.5em;\n"
            "    font-size: 0.82rem;\n"
            "    background: rgba(255,255,255,0.84);\n"
            "    border-radius: 14px;\n"
            "    overflow: hidden;\n"
            "  }\n"
            "  th {\n"
            "    background: #0f437a;\n"
            "    color: #ffffff;\n"
            "    text-align: left;\n"
            "    padding: 10px 12px;\n"
            "  }\n"
            "  td {\n"
            "    padding: 10px 12px;\n"
            "    border-bottom: 1px solid rgba(15, 67, 122, 0.10);\n"
            "    vertical-align: top;\n"
            "  }\n"
            "  tr:last-child td {\n"
            "    border-bottom: 0;\n"
            "  }\n"
            "  code {\n"
            "    font-size: 0.82rem;\n"
            "    background: rgba(15, 67, 122, 0.08);\n"
            "    padding: 0.16em 0.34em;\n"
            "    border-radius: 6px;\n"
            "  }\n"
            "  hr {\n"
            "    border: 0;\n"
            "    height: 1px;\n"
            "    background: rgba(15, 67, 122, 0.18);\n"
            "  }\n"
            "  section.lead {\n"
            "    background: radial-gradient(circle at top left, rgba(99, 179, 237, 0.30), transparent 32%), linear-gradient(135deg, #0d1b2a 0%, #16324f 55%, #23486b 100%);\n"
            "    color: #f5fbff;\n"
            "    display: flex;\n"
            "    flex-direction: column;\n"
            "    justify-content: center;\n"
            "  }\n"
            "  section.lead h1, section.lead h2, section.lead strong, section.lead blockquote {\n"
            "    color: #ffffff;\n"
            "  }\n"
            "  section.lead p, section.lead li, section.lead::after {\n"
            "    color: rgba(255,255,255,0.88);\n"
            "  }\n"
            "  section.lead h1 {\n"
            "    font-size: 2.7rem;\n"
            "    max-width: 10.5em;\n"
            "    margin-bottom: 0.28em;\n"
            "  }\n"
            "  section.lead h2 {\n"
            "    font-size: 1.14rem;\n"
            "    font-weight: 500;\n"
            "    margin-bottom: 1em;\n"
            "    max-width: 30em;\n"
            "  }\n"
            "  section.lead blockquote {\n"
            "    background: rgba(255,255,255,0.08);\n"
            "    border-left-color: rgba(255,255,255,0.7);\n"
            "    max-width: 34em;\n"
            "  }\n"
            "  section.agenda ul {\n"
            "    display: grid;\n"
            "    grid-template-columns: 1fr 1fr;\n"
            "    gap: 0.55em 1.8em;\n"
            "    padding-left: 0;\n"
            "    list-style: none;\n"
            "    margin-top: 1.2em;\n"
            "  }\n"
            "  section.agenda li {\n"
            "    background: rgba(255,255,255,0.88);\n"
            "    border: 1px solid rgba(17, 55, 99, 0.08);\n"
            "    border-radius: 14px;\n"
            "    padding: 0.8em 0.95em;\n"
            "    box-shadow: 0 12px 28px rgba(17, 55, 99, 0.08);\n"
            "    margin: 0;\n"
            "  }\n"
            "  section.metrics .metrics-grid {\n"
            "    display: grid;\n"
            "    grid-template-columns: 1fr 1fr;\n"
            "    gap: 0.85em;\n"
            "    margin-top: 1.1em;\n"
            "  }\n"
            "  section.metrics .metric-card {\n"
            "    background: rgba(255,255,255,0.92);\n"
            "    border: 1px solid rgba(16, 36, 61, 0.08);\n"
            "    border-radius: 16px;\n"
            "    padding: 0.95em 1.05em;\n"
            "    box-shadow: 0 16px 30px rgba(23, 45, 78, 0.08);\n"
            "  }\n"
            "  section.metrics .metric-label {\n"
            "    display: block;\n"
            "    font-size: 0.78rem;\n"
            "    font-weight: 700;\n"
            "    letter-spacing: 0.04em;\n"
            "    text-transform: uppercase;\n"
            "    color: #55708f;\n"
            "    margin-bottom: 0.4em;\n"
            "  }\n"
            "  section.metrics .metric-value {\n"
            "    display: block;\n"
            "    font-size: 1.22rem;\n"
            "    font-weight: 800;\n"
            "    color: #10243d;\n"
            "    margin-bottom: 0.28em;\n"
            "  }\n"
            "  section.metrics .metric-text {\n"
            "    display: block;\n"
            "    font-size: 0.83rem;\n"
            "    line-height: 1.34;\n"
            "    color: #455d79;\n"
            "  }\n"
            "  section.timeline ol {\n"
            "    list-style: none;\n"
            "    counter-reset: steps;\n"
            "    margin-top: 1em;\n"
            "    padding-left: 0;\n"
            "  }\n"
            "  section.timeline li {\n"
            "    counter-increment: steps;\n"
            "    position: relative;\n"
            "    padding: 0.18em 0 0.9em 3em;\n"
            "    margin: 0;\n"
            "  }\n"
            "  section.timeline li::before {\n"
            "    content: counter(steps);\n"
            "    position: absolute;\n"
            "    left: 0;\n"
            "    top: 0.06em;\n"
            "    width: 2em;\n"
            "    height: 2em;\n"
            "    border-radius: 999px;\n"
            "    background: #163f73;\n"
            "    color: #fff;\n"
            "    display: grid;\n"
            "    place-items: center;\n"
            "    font-size: 0.75rem;\n"
            "    font-weight: 800;\n"
            "  }\n"
            "  section.timeline li::after {\n"
            "    content: '';\n"
            "    position: absolute;\n"
            "    left: 0.95em;\n"
            "    top: 2.2em;\n"
            "    bottom: -0.1em;\n"
            "    width: 2px;\n"
            "    background: rgba(22, 63, 115, 0.16);\n"
            "  }\n"
            "  section.timeline li:last-child::after {\n"
            "    display: none;\n"
            "  }\n"
            "  section.compare table {\n"
            "    margin-top: 1.1em;\n"
            "    box-shadow: 0 16px 32px rgba(23, 45, 78, 0.08);\n"
            "  }\n"
            "  section.closing {\n"
            "    background: linear-gradient(135deg, #10243d 0%, #1c4f88 100%);\n"
            "    color: #f4f8fc;\n"
            "    display: flex;\n"
            "    flex-direction: column;\n"
            "    justify-content: center;\n"
            "  }\n"
            "  section.closing h1, section.closing h2, section.closing strong {\n"
            "    color: #ffffff;\n"
            "  }\n"
            "  section.closing p, section.closing li, section.closing::after {\n"
            "    color: rgba(255,255,255,0.9);\n"
            "  }\n"
            "  section.closing ul {\n"
            "    margin-top: 1em;\n"
            "  }\n"
            "---\n\n"
        )
        return frontmatter + "\n\n---\n\n".join(slides).strip() + "\n"

    def _split_raw_slides(self, text: str) -> list[str]:
        normalized = self._normalize_slide_source(text)
        chunks = [self._sanitize_slide_chunk(chunk) for chunk in re.split(r"(?m)^\s*---+\s*$", normalized)]
        return [chunk for chunk in chunks if chunk]

    def _looks_like_marp_deck(self, text: str) -> bool:
        if "marp: true" in text.lower():
            return True
        return text.count("\n---") >= 1 and ("# " in text or "## " in text)

    def _split_into_slides(self, text: str) -> list[str]:
        normalized = self._normalize_slide_source(text)
        if "\n---" in normalized:
            parts = [self._sanitize_slide_chunk(chunk) for chunk in re.split(r"(?m)^\s*---+\s*$", normalized) if chunk.strip()]
            if parts:
                return parts
        sections = []
        current: list[str] = []
        for line in normalized.splitlines():
            stripped = line.strip()
            if re.match(r"^#{1,2}\s+", stripped) and current:
                sections.append(self._sanitize_slide_chunk("\n".join(current)))
                current = [stripped]
            else:
                current.append(stripped)
        if current:
            sections.append(self._sanitize_slide_chunk("\n".join(current)))
        sections = [section for section in sections if section]
        if len(sections) <= 1:
            return self._synthesize_slide_deck(normalized)
        normalized = [self._normalize_slide(section) for section in sections]
        if normalized:
            normalized[0] = self._upgrade_cover_slide(normalized[0])
        return normalized

    def _normalize_slide_source(self, text: str) -> str:
        cleaned = str(text or "").replace("\r\n", "\n").replace("\r", "\n").strip()
        if not cleaned:
            return ""
        cleaned = re.sub(r"^```(?:markdown|md|text|html)?\s*", "", cleaned, flags=re.I)
        cleaned = re.sub(r"\s*```$", "", cleaned)
        lines = cleaned.splitlines()
        if lines and lines[0].strip() == "---":
            for idx in range(1, min(len(lines), 24)):
                if lines[idx].strip() == "---":
                    frontmatter_body = lines[1:idx]
                    if any(":" in item for item in frontmatter_body):
                        lines = lines[idx + 1 :]
                    break
        cleaned = "\n".join(lines)
        cleaned = re.sub(r"(?m)^\s*>\s*", "", cleaned)
        cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
        return cleaned.strip()

    def _sanitize_slide_chunk(self, chunk: str) -> str:
        lines: list[str] = []
        for raw_line in chunk.splitlines():
            cleaned = self._clean_slide_line(raw_line)
            if cleaned:
                lines.append(cleaned)
        return "\n".join(lines).strip()

    def _clean_slide_line(self, raw_line: str) -> str:
        stripped = raw_line.strip()
        if not stripped:
            return ""
        if stripped.startswith("<!--") or stripped.startswith("```"):
            return ""
        if re.fullmatch(r"-{3,}", stripped):
            return ""
        if re.match(r"^(marp|theme|paginate|size|headingDivider|footer)\s*:", stripped, flags=re.I):
            return ""
        if stripped.lower() == "style:" or re.match(r"^[a-z-]+\s*\{$", stripped.lower()):
            return ""
        if stripped.startswith("![") and "](" in stripped:
            return ""
        bullet_prefix = ""
        if re.match(r"^[-*]\s+", stripped):
            bullet_prefix = "- "
            stripped = re.sub(r"^[-*]\s+", "", stripped)
        elif re.match(r"^\d+[.)]\s+", stripped):
            bullet_prefix = "- "
            stripped = re.sub(r"^\d+[.)]\s+", "", stripped)
        stripped = re.sub(r"^#{1,6}\s*", "", stripped)
        stripped = re.sub(r"^>\s*", "", stripped)
        stripped = self._clean_inline_markdown(stripped)
        if not stripped or stripped in {"---", "--"}:
            return ""
        if re.fullmatch(r"[_*`~#\\\-.|: ]+", stripped):
            return ""
        return f"{bullet_prefix}{stripped}" if bullet_prefix else stripped

    def _synthesize_slide_deck(self, text: str) -> list[str]:
        title = self._slide_title_from_text(text)
        sentences = [chunk.strip(" -") for chunk in re.split(r"(?<=[\.\!\?])\s+", " ".join(text.split())) if chunk.strip()]
        bullets = [self._clean_inline_markdown(chunk) for chunk in sentences if chunk.strip()][:12]
        if not bullets:
            bullets = ["Contexto principal", "Pontos-chave", "Recomendacao final"]
        grouped = [bullets[index : index + 3] for index in range(0, len(bullets), 3)]
        slides = [self._upgrade_cover_slide(f"# {title}\n\n## {bullets[0]}")]
        for index, group in enumerate(grouped[:3], start=1):
            heading = "## Destaques" if index == 1 else f"## Topico {index}"
            slides.append("\n".join([heading, "", *[f"- {item}" for item in group]]))
        slides.append(self._render_closing_slide(f"## Fechamento\n\n- {bullets[-1]}", title))
        return [self._sanitize_slide_chunk(slide) for slide in slides if slide.strip()]

    def _needs_slide_fallback(self, source_text: str, user_request: str) -> bool:
        compact = " ".join(source_text.split()).lower()
        if len(compact) < 260:
            return True
        weak_markers = [
            "documento executivo",
            "contexto principal",
            "pontos-chave",
            "recomendacao final",
            "recomendação final",
            "panorama claro, visual e pronto",
            "estrutura base de documento",
            "este documento apresenta",
            "contexto geral",
            "em um documento final",
            "este pdf e docx",
        ]
        if any(marker in compact for marker in weak_markers):
            return True
        topic = self._extract_topic(user_request).lower()
        topic_terms = [term for term in re.split(r"\W+", topic) if len(term) >= 5][:4]
        if topic_terms and not any(term in compact for term in topic_terms):
            return True
        return False

    def _fallback_slide_deck_text(self, user_request: str, title: str) -> str:
        topic = self._clean_inline_markdown(self._extract_topic(user_request) or title)
        lowered = user_request.lower()
        direction = self._pptx_creative_direction(user_request, title)
        return self._adaptive_fallback_slide_deck_text(topic, user_request, direction)
        if direction.domain == "food":
            return f"""# {topic}
## Sabor, desejo e decisao em uma historia simples

---

## Apetite imediato
- **Primeira impressao** - promessa clara antes do preco
- **Desejo visual** - ingrediente, textura e frescor em foco
- **Escolha facil** - poucas opcoes, decisao rapida

---

## Experiencia servida
- **Produto estrela** - destaque para o item mais memoravel
- **Prova sensorial** - aroma, crocancia, cremosidade ou frescor
- **Momento de consumo** - almoco, encontro, presente ou pausa
- **Chamada direta** - pedir, reservar ou experimentar

---

## Cardapio com hierarquia
- **Entrada** - despertar curiosidade sem cansar
- **Principal** - concentrar valor no prato assinatura
- **Complemento** - aumentar ticket sem poluir a escolha
- **Finalizacao** - sobremesa, bebida ou combo

---

## Oferta que converte
- **Foto honesta** - alimento reconhecivel e apetitoso
- **Beneficio claro** - sabor, praticidade ou exclusividade
- **Preco legivel** - sem esconder a decisao
- **Urgencia leve** - horario, estoque ou promocao

---

## Lancamento em etapas
- **Preparar** - selecionar prato, foto e promessa
- **Testar** - validar mensagem com publico pequeno
- **Divulgar** - canais, horarios e criativos
- **Ajustar** - medir pedidos e repetir vencedores

---

## Decisao recomendada
- **Hoje** - escolher o produto heroi
- **7 dias** - produzir imagens e oferta
- **30 dias** - medir conversao e ticket medio
"""
        if direction.domain == "animals":
            return f"""# {topic}
## Observacao, cuidado e contexto natural

---

## O que observar
- **Comportamento** - sinais que revelam adaptacao
- **Habitat** - ambiente molda rotina e sobrevivencia
- **Interacao** - relacao com grupo, alimento e territorio

---

## Leitura do animal
- **Aparencia** - marcas, porte e diferencas visiveis
- **Rotina** - sono, busca por alimento e deslocamento
- **Defesa** - fuga, camuflagem ou protecao
- **Vinculo** - cuidado parental, bando ou independencia

---

## Habitat em foco
- **Recursos** - alimento, agua e abrigo disponiveis
- **Riscos** - predadores, clima e pressao humana
- **Adaptacoes** - corpo e comportamento trabalhando juntos
- **Equilibrio** - papel na cadeia e no ecossistema

---

## Cuidado responsavel
- **Bem-estar** - espaco, estimulo e seguranca
- **Saude** - prevencao, sinais e acompanhamento
- **Manejo** - rotina coerente com a especie
- **Educacao** - reduzir medo e aproximar com respeito

---

## Aprendizado em campo
- **Identificar** - especie, ambiente e sinais
- **Registrar** - fotos, notas e padroes
- **Interpretar** - comportamento ligado ao contexto
- **Proteger** - acao adequada ao risco observado

---

## Decisao recomendada
- **Agora** - definir objetivo da observacao
- **Proximo passo** - coletar evidencias visuais
- **Continuo** - agir sem romper o habitat
"""
        if direction.domain == "education":
            return f"""# {topic}
## Aprender melhor com percurso claro

---

## Ponto de partida
- **Objetivo** - uma competencia por vez
- **Contexto** - por que isso importa agora
- **Aplicacao** - onde o aluno usa o conhecimento

---

## Trilha de aula
- **Conceito** - ideia central em linguagem simples
- **Exemplo** - caso concreto para fixar
- **Pratica** - exercicio curto e verificavel
- **Feedback** - correcao objetiva e proximo passo

---

## Metodo de aprendizagem
- **Atencao** - foco visual sem excesso
- **Memoria** - repeticao com variacao
- **Autonomia** - pequenas decisoes do aluno
- **Evidencia** - prova do que foi entendido

---

## Material que ajuda
- **Resumo** - pontos essenciais sem paragrafo longo
- **Atividade** - tarefa conectada ao objetivo
- **Rubrica** - criterio claro de avaliacao
- **Revisao** - volta rapida ao que falhou

---

## Plano em etapas
- **Diagnosticar** - nivel e duvidas reais
- **Explicar** - conceito com exemplo visual
- **Praticar** - desafio progressivo
- **Consolidar** - revisao e criterio de dominio

---

## Decisao recomendada
- **Hoje** - escolher a competencia central
- **Aula** - alternar explicacao e pratica
- **Depois** - revisar pelo erro mais comum
"""
        if direction.domain == "health":
            return f"""# {topic}
## Clareza, cuidado e decisao segura

---

## Necessidade principal
- **Sintoma ou risco** - entender o ponto de atencao
- **Contexto** - rotina, historico e fatores associados
- **Conduta** - orientar sem gerar confusao

---

## Leitura do cuidado
- **Prevencao** - agir antes da complicacao
- **Acompanhamento** - sinais, progresso e retorno
- **Adesao** - orientacao facil de seguir
- **Seguranca** - limites e quando procurar ajuda

---

## Jornada do paciente
- **Escuta** - queixa e objetivo real
- **Avaliacao** - sinais e evidencias
- **Plano** - passos claros e viaveis
- **Retorno** - ajuste com base em resposta

---

## Indicadores de qualidade
- **Entendimento** - paciente sabe o que fazer
- **Continuidade** - cuidado nao para na consulta
- **Risco** - sinais de alerta comunicados
- **Resultado** - evolucao acompanhada com criterio

---

## Plano pratico
- **Orientar** - linguagem simples e direta
- **Registrar** - informacoes essenciais
- **Monitorar** - sinais de melhora ou alerta
- **Reavaliar** - decidir proximo passo

---

## Decisao recomendada
- **Agora** - separar orientacao de alerta
- **7 dias** - acompanhar resposta
- **30 dias** - revisar plano e adesao
"""
        if direction.domain == "fashion":
            return f"""# {topic}
## Estilo, desejo e identidade visual

---

## Impressao imediata
- **Silhueta** - forma comunica antes do detalhe
- **Textura** - material cria percepcao de valor
- **Contraste** - cor guia o olhar

---

## Narrativa da peca
- **Ocasião** - quando e por que usar
- **Atitude** - casual, sofisticada ou ousada
- **Combinacao** - peca principal e apoio
- **Desejo** - detalhe que torna memoravel

---

## Direcao visual
- **Modelo** - postura alinhada ao publico
- **Luz** - pele, tecido e volume favorecidos
- **Cenario** - contexto sem competir com a peca
- **Edicao** - recorte limpo e foco no produto

---

## Colecao com hierarquia
- **Heroi** - peca assinatura
- **Base** - itens de rotacao
- **Acessorio** - elevar composicao
- **Campanha** - frase curta e imagem forte

---

## Lancamento em etapas
- **Curadoria** - selecionar pecas-chave
- **Producao** - foto, styling e promessa
- **Publicacao** - canais e sequencia
- **Aprendizado** - medir interesse e conversao

---

## Decisao recomendada
- **Hoje** - definir peca hero
- **Semana** - produzir editorial enxuto
- **Mes** - repetir visual vencedor
"""
        wants_iso = any(token in lowered for token in ["iso", "norma", "normas", "certificacao", "certifica"])
        wants_steps = any(token in lowered for token in ["roadmap", "passos", "plano", "implementacao"])
        standards = [
            "ISO 55001 - governanca do ciclo de vida",
            "ISO/IEC 19770 - inventario, licencas e conformidade",
            "ISO/IEC 27001 - seguranca aplicada a ativos criticos",
            "ISO/IEC 20000-1 - servicos conectados ao catalogo",
        ]
        standards_slide = "\n".join(f"- {item}" for item in standards)
        general_references = "\n".join(
            [
                "- Ciclo de vida - entrada, uso, renovacao e descarte",
                "- Controles criticos - dono, risco e evidencias",
                "- Governanca operacional - regras, excecoes e auditoria",
                "- Metricas executivas - custo, risco e disponibilidade",
            ]
        )
        references_slide = standards_slide if wants_iso else general_references
        roadmap_title = "Execucao em 90 dias" if wants_steps else "Modelo de controle"
        return f"""# {topic}
## Controle, risco e valor em uma base unica

---

## Risco invisivel
- **Inventario falho** - ativos sem dono nem status
- **Custo disperso** - licencas e contratos sem reconciliacao
- **Seguranca exposta** - criticidade fora do radar

---

## Prioridades executivas
- **Visibilidade** - inventario unico e confiavel
- **Controle** - custo, contrato e uso conectados
- **Governanca** - dono claro para cada ativo
- **Seguranca** - risco priorizado por impacto

---

## Referencias que sustentam
{references_slide}

---

## Arquitetura de controle
- **Base unica** - inventario, contratos e telemetria
- **Dono do ativo** - responsabilidade por custo e risco
- **Ciclo de vida** - entrada, uso, renovacao e descarte
- **Indicadores** - cobertura, conformidade e disponibilidade

---

## {roadmap_title}
- **Diagnosticar** - fontes, lacunas e ativos criticos
- **Integrar** - inventario, descoberta e contratos
- **Governar** - papeis, politicas e evidencias
- **Otimizar** - licencas, risco e ativos ociosos

---

## Indicadores que importam
- **Cobertura** - ativos classificados e com dono
- **Economia** - desperdicio encontrado e removido
- **Risco** - ativos criticos com plano ativo
- **Conformidade** - evidencias prontas para auditoria

---

## Decisao recomendada
- **30 dias** - consolidar fontes confiaveis
- **60 dias** - priorizar custo e risco alto
- **90 dias** - institucionalizar governanca recorrente
"""

    def _adaptive_fallback_slide_deck_text(self, topic: str, user_request: str, direction: DeckVisualDirection) -> str:
        topic = self._clean_inline_markdown(topic or "Tema principal").strip(" .:-")
        topic_l = topic.lower()
        lowered = self._normalize_for_match(user_request)
        wants_launch = any(token in lowered for token in ["anunciar", "divulgar", "vender", "promover", "lancar", "lançamento", "campanha"])
        wants_teach = any(token in lowered for token in ["aula", "curso", "didatico", "ensinar", "explicar", "treinamento"])
        wants_plan = any(token in lowered for token in ["plano", "roadmap", "passos", "implementar", "estrategia"])

        if wants_launch:
            promise = "uma campanha com desejo, clareza e acao"
            tension = "o publico decide rapido quando a promessa fica concreta"
            objective = "transformar atencao em interesse e interesse em acao"
            closing = "publicar a versao mais forte, medir resposta e ajustar"
            steps = ["Escolher gancho", "Criar prova visual", "Publicar oferta", "Medir resposta"]
        elif wants_teach:
            promise = "uma explicacao clara, memoravel e facil de seguir"
            tension = "o aluno perde o fio quando tudo parece ter o mesmo peso"
            objective = "organizar o tema em etapas, exemplos e pratica"
            closing = "validar compreensao e reforcar o ponto mais dificil"
            steps = ["Abrir contexto", "Explicar conceito", "Mostrar exemplo", "Fixar pratica"]
        elif wants_plan:
            promise = "um caminho objetivo para sair da ideia e chegar na execucao"
            tension = "sem sequencia clara, o tema vira lista solta"
            objective = "separar prioridade, responsavel, evidencia e proximo passo"
            closing = "começar pequeno, medir sinal e evoluir a partir do resultado"
            steps = ["Diagnosticar", "Priorizar", "Executar", "Aprender"]
        else:
            promise = "uma leitura visual feita para entender, decidir e lembrar"
            tension = "o tema perde forca quando vira resumo generico"
            objective = "dar forma ao assunto com contexto, contraste e exemplos"
            closing = "transformar a leitura em uma decisao ou proximo passo"
            steps = ["Contexto", "Contraste", "Evidencia", "Proximo passo"]

        domain_angles = {
            "food": ["sensorial", "produto", "experiencia", "oferta"],
            "animals": ["habitat", "comportamento", "adaptacao", "cuidado"],
            "fashion": ["silhueta", "textura", "identidade", "desejo"],
            "health": ["cuidado", "prevencao", "seguranca", "continuidade"],
            "education": ["conceito", "exemplo", "pratica", "retencao"],
            "technology": ["sistema", "dados", "risco", "automacao"],
            "business": ["valor", "risco", "prioridade", "decisao"],
        }
        angles = domain_angles.get(direction.domain, ["contexto", "valor", "prova", "acao"])
        title_label = self._compact_title(topic, 8)
        return f"""# {title_label}
## {promise}

---

## Por que importa
- **Tensao** - {tension}
- **Intencao** - {objective}
- **Leitura** - {direction.mood}

---

## O que destacar
- **{angles[0].title()}** - ponto que abre a percepcao sobre {topic_l}
- **{angles[1].title()}** - detalhe que torna o assunto especifico
- **{angles[2].title()}** - evidencia que evita aparencia generica
- **{angles[3].title()}** - acao esperada depois da leitura

---

## Narrativa visual
- **Clima** - visual alinhado ao universo de {topic_l}
- **Foco** - uma ideia dominante por slide
- **Ritmo** - alternar impacto, detalhe e decisao
- **Respiro** - menos texto, mais hierarquia

---

## Sequencia recomendada
- **{steps[0]}** - abrir com o ponto mais reconhecivel
- **{steps[1]}** - mostrar o contraste que muda a leitura
- **{steps[2]}** - trazer prova, exemplo ou oferta
- **{steps[3]}** - encerrar com acao objetiva

---

## Criterios de qualidade
- **Pertinencia** - parecer feito para {topic_l}
- **Clareza** - titulo forte e texto curto
- **Identidade** - cor, imagem e tom coerentes
- **Autonomia** - respeitar publico, canal e estilo pedidos

---

## Proximo passo
- **Agora** - revisar a intencao principal
- **Depois** - trocar exemplos pelo contexto real do usuario
- **Final** - {closing}
"""

    def _polish_slide_deck(self, slides: list[str], title: str, slide_visuals: dict[int, str]) -> list[str]:
        polished: list[str] = []
        total = len(slides)
        for index, slide in enumerate(slides):
            working = slide.strip()
            slide_class = self._infer_slide_class(working, index, total)
            if slide_class == "lead":
                working = self._upgrade_cover_slide(working)
            elif slide_class == "agenda":
                working = self._normalize_agenda_slide(working)
            elif slide_class == "metrics":
                working = self._render_metrics_slide(working)
            elif slide_class == "timeline":
                working = self._render_timeline_slide(working)
            elif slide_class == "closing":
                working = self._render_closing_slide(working, title)
            visual_path = slide_visuals.get(index)
            if visual_path:
                working = self._inject_visual_asset(working, visual_path, slide_class)
            working = self._inject_slide_class(working, slide_class)
            polished.append(working)
        return polished

    def _generate_slide_visuals(
        self,
        title: str,
        user_request: str,
        slides: list[str],
        folder: Path,
    ) -> dict[int, str]:
        if not slides or not self._wants_ai_slide_images(user_request):
            return {}
        selected_indexes = self._select_visual_slides(slides, user_request)
        if not selected_indexes:
            return {}
        assets_dir = folder / "assets"
        assets_dir.mkdir(parents=True, exist_ok=True)
        visuals: dict[int, str] = {}
        for slide_index in selected_indexes:
            prompt = self._slide_image_prompt(title, user_request, slides[slide_index], slide_index, len(slides))
            output_path = assets_dir / f"slide-{slide_index + 1:02d}-visual.png"
            try:
                if self._download_wikimedia_reference_image(title, slides[slide_index], output_path):
                    visuals[slide_index] = f"assets/{output_path.name}"
                    continue
                if self.settings.pollinations_api_key:
                    self._download_pollinations_image(prompt, output_path)
                    visuals[slide_index] = f"assets/{output_path.name}"
            except Exception:
                continue
        return visuals

    def _select_visual_slides(self, slides: list[str], user_request: str = "") -> list[int]:
        choices: list[int] = []
        lowered_request = user_request.lower()
        if slides:
            choices.append(0)
        for idx, slide in enumerate(slides[1:-1], start=1):
            lower = slide.lower()
            if any(token in lower for token in ["impacto", "cenario", "cenario", "mercado", "processo", "cronograma", "arquitetura", "roadmap", "estrategia", "estratégia"]):
                choices.append(idx)
            if len(choices) >= 4:
                break
        if len(choices) < 2 and len(slides) > 2:
            choices.append(1)
        if any(token in lowered_request for token in ["imagem", "imagens", "visual", "ilustrado", "com fotos", "com foto"]) and len(slides) > 3:
            choices.append(min(2, len(slides) - 2))
        deduped: list[int] = []
        for item in choices:
            if item not in deduped:
                deduped.append(item)
        return deduped[:4]

    def _wants_ai_slide_images(self, user_request: str) -> bool:
        lowered = user_request.lower()
        return any(
            token in lowered
            for token in [
                "com imagem",
                "com imagens",
                "com foto",
                "com fotos",
                "ilustrado",
                "ilustrações",
                "ilustracoes",
                "visual com imagem",
                "use imagens",
                "usar imagens",
            ]
        )

    def _slide_image_prompt(
        self,
        deck_title: str,
        user_request: str,
        slide: str,
        slide_index: int,
        total_slides: int,
    ) -> str:
        lines = [line.strip() for line in slide.splitlines() if line.strip() and not line.strip().startswith("<!--")]
        heading = re.sub(r"^#{1,3}\s*", "", lines[0]).strip() if lines else deck_title
        bullets = [
            re.sub(r"^[-*]\s+", "", line).strip()
            for line in lines[1:]
            if re.match(r"^[-*]\s+", line)
        ][:3]
        narrative = "; ".join(bullets)[:360]
        direction = self._pptx_creative_direction(user_request, deck_title)
        style_hint = self._image_style_hint(user_request, direction)
        is_cover = slide_index == 0
        if is_cover:
            return (
                f"Premium presentation cover image for '{deck_title}'. "
                f"Theme: {heading}. "
                f"Editorial, polished storytelling, {direction.mood}, "
                f"{style_hint} clean composition with negative space for title text, subtle depth, modern lighting, no text, no watermark. "
                f"Context: {user_request[:260]}"
            )
        return (
            f"Presentation visual for slide {slide_index + 1} of {total_slides} about '{deck_title}'. "
            f"Slide topic: {heading}. "
            f"Key points: {narrative or user_request[:220]}. "
            f"Subject-specific visual direction: {direction.mood}. "
            "Professional editorial illustration or photoreal concept for a premium presentation, "
            f"{style_hint} clean composition, sophisticated color palette, suitable for split-slide layout, no text, no watermark."
        )

    def _image_style_hint(self, user_request: str, direction: DeckVisualDirection | None = None) -> str:
        lowered = user_request.lower()
        if any(token in lowered for token in ["realista", "fotorealista", "foto"]):
            return f"Photoreal, realistic, {direction.image_style if direction else 'premium editorial photography'},"
        if any(token in lowered for token in ["3d", "futurista", "neon"]):
            return f"Futuristic 3D render language, {direction.image_style if direction else 'premium editorial composition'},"
        if any(token in lowered for token in ["minimal", "minimalista", "clean"]):
            return f"Minimal editorial visual language, {direction.image_style if direction else 'premium subject-specific composition'},"
        if any(token in lowered for token in ["luxo", "premium", "executivo"]):
            return f"Luxury executive editorial visual language, {direction.image_style if direction else 'premium composition'},"
        return f"{direction.image_style if direction else 'Professional presentation visual language'},"

    def _download_wikimedia_reference_image(self, deck_title: str, raw_slide: str, output_path: Path) -> bool:
        terms = self._wikimedia_query_terms(deck_title, raw_slide)
        if not terms:
            return False
        api_url = "https://commons.wikimedia.org/w/api.php"
        with httpx.Client(timeout=30) as client:
            for term in terms:
                params = {
                    "action": "query",
                    "format": "json",
                    "generator": "search",
                    "gsrsearch": term,
                    "gsrnamespace": 6,
                    "gsrlimit": 6,
                    "prop": "imageinfo",
                    "iiprop": "url",
                    "iiurlwidth": 1920,
                    "origin": "*",
                }
                response = client.get(api_url, params=params)
                if response.status_code >= 400:
                    continue
                payload = response.json()
                pages = (payload.get("query") or {}).get("pages") or {}
                for page in pages.values():
                    title = str(page.get("title") or "").lower()
                    if any(token in title for token in ["logo", "icon", "flag", "coat of arms"]):
                        continue
                    info = (page.get("imageinfo") or [{}])[0]
                    image_url = info.get("thumburl") or info.get("url")
                    if not image_url or not re.search(r"\.(jpg|jpeg|png)(?:\?|$)", image_url, flags=re.I):
                        continue
                    image_response = client.get(image_url, timeout=45)
                    if image_response.status_code >= 400:
                        continue
                    content = image_response.content
                    if len(content) < 12_000:
                        continue
                    output_path.write_bytes(content)
                    return True
        return False

    def _wikimedia_query_terms(self, deck_title: str, raw_slide: str) -> list[str]:
        lines = [self._clean_inline_markdown(line) for line in raw_slide.splitlines() if line.strip()]
        heading = ""
        for line in lines:
            if line.startswith("#"):
                heading = re.sub(r"^#{1,4}\s*", "", line).strip()
                break
        if not heading and lines:
            heading = lines[0]
        heading = re.sub(r"[^A-Za-z0-9\s-]", " ", heading).strip()
        title = re.sub(r"[^A-Za-z0-9\s-]", " ", deck_title).strip()
        candidates = [
            heading,
            f"{heading} business presentation" if heading else "",
            title,
            f"{title} corporate" if title else "",
        ]
        terms: list[str] = []
        for candidate in candidates:
            normalized = " ".join(candidate.split())
            if len(normalized) < 4:
                continue
            if normalized not in terms:
                terms.append(normalized)
        return terms[:4]

    def _download_pollinations_image(self, prompt: str, output_path: Path) -> None:
        headers = {"Authorization": f"Bearer {self.settings.pollinations_api_key}"}
        payload = {
            "model": self.settings.pollinations_image_model,
            "prompt": prompt,
            "size": "1536x1024",
            "quality": self.settings.pollinations_image_quality,
            "response_format": "b64_json",
        }
        with httpx.Client(timeout=150) as client:
            response = client.post("https://gen.pollinations.ai/v1/images/generations", headers=headers, json=payload)
            response.raise_for_status()
            data = response.json()
        image_payload = (data.get("data") or [{}])[0]
        image_b64 = image_payload.get("b64_json")
        image_url = image_payload.get("url")
        if image_b64:
            output_path.write_bytes(base64.b64decode(image_b64))
            return
        if image_url:
            with httpx.Client(timeout=150) as client:
                response = client.get(image_url)
                response.raise_for_status()
            output_path.write_bytes(response.content)
            return
        raise RuntimeError("Imagem de slide nao retornou conteudo utilizavel.")

    def _inject_visual_asset(self, slide: str, visual_path: str, slide_class: str) -> str:
        normalized = visual_path.replace("\\", "/")
        if slide_class == "lead":
            visual_line = f"![bg right:42%]({quote(normalized, safe='/:.-_')})"
        else:
            visual_line = f"![bg right:36%]({quote(normalized, safe='/:.-_')})"
        return f"{visual_line}\n\n{slide.strip()}"

    def _build_slide_models(self, raw_slides: list[str], deck_title: str, slide_visuals: dict[int, str]) -> list[Slide]:
        slides: list[Slide] = []
        total = len(raw_slides)
        for index, raw_slide in enumerate(raw_slides):
            layout = self._infer_slide_class(raw_slide, index, total)
            lines = [line for line in (self._clean_slide_line(item) for item in raw_slide.splitlines()) if line]
            title_line = next((line for line in lines if not line.startswith("- ")), lines[0] if lines else deck_title)
            title = self._humanize_slide_title(title_line, index, deck_title)
            bullets = [self._clean_inline_markdown(re.sub(r"^-\s+", "", line).strip()) for line in lines if line.startswith("- ")]
            extras = [line for line in lines if line != title_line and not line.startswith("- ")]
            rows = self._extract_slide_table(raw_slide)
            kicker = self._slide_kicker(deck_title, layout, index)
            body = [self._clean_inline_markdown(item) for item in extras if not self._looks_like_table_line(item)]
            body = [item for item in body if item and item != title and item.lower() != kicker.lower()]
            bullets = [item for item in bullets if item and item != title][:4]
            if layout == "lead" and not body and bullets:
                body = [bullets[0]]
                bullets = bullets[1:]
            if layout == "lead":
                bullets = bullets[:2]
            slides.append(
                Slide(
                    title=title,
                    kicker=kicker,
                    body=body[:2],
                    bullets=bullets,
                    rows=rows,
                    layout=layout,
                    visual=slide_visuals.get(index),
                )
            )
        return slides

    def _build_presentation_pptx(self, path: Path, deck_title: str, slides: list[Slide], folder: Path, user_request: str = "") -> str:
        presentation = Presentation()
        presentation.slide_width = Inches(13.333)
        presentation.slide_height = Inches(7.5)
        blank_layout = presentation.slide_layouts[6]
        total = len(slides)
        direction = self._pptx_creative_direction(user_request, deck_title)
        for index, slide in enumerate(slides):
            ppt_slide = presentation.slides.add_slide(blank_layout)
            self._render_pptx_slide(ppt_slide, slide, deck_title, index, total, folder, direction)
        presentation.save(str(path))
        return "python-pptx"

    def _build_presentation_pdf_from_pptx(self, pptx_path: Path, pdf_path: Path, slides: list[Slide]) -> str:
        office_binary = shutil.which("soffice") or shutil.which("libreoffice")
        if office_binary:
            try:
                subprocess.run(
                    [
                        office_binary,
                        "--headless",
                        "--convert-to",
                        "pdf",
                        "--outdir",
                        str(pdf_path.parent),
                        str(pptx_path),
                    ],
                    check=True,
                    capture_output=True,
                    text=True,
                    timeout=180,
                )
                generated = pdf_path.parent / f"{pptx_path.stem}.pdf"
                if generated.exists() and generated != pdf_path:
                    generated.replace(pdf_path)
                if pdf_path.exists():
                    return "libreoffice-pptx"
            except Exception:
                pass
        self._build_slide_pdf_with_reportlab(pdf_path, slides)
        return "reportlab-slides"

    def _render_pptx_slide(
        self,
        ppt_slide: Any,
        slide: Slide,
        deck_title: str,
        index: int,
        total: int,
        folder: Path,
        direction: DeckVisualDirection,
    ) -> None:
        if slide.layout == "lead":
            self._render_pptx_cover(ppt_slide, slide, deck_title, index, total, folder, direction)
            return
        if slide.layout == "closing":
            self._render_pptx_closing(ppt_slide, slide, deck_title, index, total, folder, direction)
            return
        self._render_pptx_content(ppt_slide, slide, deck_title, index, total, direction)

    def _render_pptx_cover(self, ppt_slide: Any, slide: Slide, deck_title: str, index: int, total: int, folder: Path, direction: DeckVisualDirection) -> None:
        bg = self._pptx_color(direction.dark_bg)
        ink = self._pptx_color("#ffffff")
        muted = self._pptx_color(direction.dark_muted)
        accent = self._pptx_color(direction.accent)
        self._pptx_add_shape(ppt_slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, 13.333, 7.5, bg)
        self._pptx_add_shape(ppt_slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, 0.12, 7.5, accent)
        self._pptx_add_shape(ppt_slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 8.7, 0, 0.012, 7.5, self._pptx_color(direction.dark_divider))
        self._pptx_add_shape(ppt_slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0.72, 0.86, 0.9, 0.045, accent)
        self._pptx_add_textbox(ppt_slide, 0.72, 0.48, 6.3, 0.26, deck_title.upper()[:74], 9.5, color=muted, bold=True)
        self._pptx_add_textbox(ppt_slide, 11.82, 0.48, 0.9, 0.24, f"{index + 1} / {total}", 10, color=muted, bold=True, align=PP_ALIGN.RIGHT)

        title = self._compact_title(slide.title, 9)
        self._pptx_add_textbox(ppt_slide, 0.72, 1.48, 7.0, 1.72, title, 34, color=ink, bold=True)
        subtitle = (slide.body or slide.bullets or ["Apresentacao executiva com narrativa clara."])[0]
        self._pptx_add_textbox(ppt_slide, 0.74, 3.34, 5.95, 0.64, self._compact_copy(subtitle, 11), 17, color=muted)

        points = self._pptx_visual_points(slide)
        self._pptx_add_textbox(ppt_slide, 9.2, 1.22, 2.6, 0.22, direction.cover_label, 9.5, color=accent, bold=True)
        for idx, point in enumerate(points[:3], start=1):
            y_pos = 1.72 + (idx - 1) * 1.22
            self._pptx_add_textbox(ppt_slide, 9.18, y_pos, 0.42, 0.24, f"{idx:02d}", 10, color=accent, bold=True)
            self._pptx_add_textbox(ppt_slide, 9.78, y_pos - 0.03, 2.82, 0.42, self._compact_copy(point, 9), 13.5, color=ink, bold=True)
            self._pptx_add_shape(ppt_slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 9.18, y_pos + 0.62, 3.1, 0.012, self._pptx_color(direction.dark_line))

        if slide.visual:
            self._pptx_try_add_image(ppt_slide, slide.visual, folder, 9.18, 5.28, 3.1, 1.24)

    def _render_pptx_content(self, ppt_slide: Any, slide: Slide, deck_title: str, index: int, total: int, direction: DeckVisualDirection) -> None:
        bg = self._pptx_color(direction.light_bg)
        ink = self._pptx_color(direction.ink)
        muted = self._pptx_color(direction.muted)
        accent = self._pptx_color(direction.accent)
        self._pptx_add_shape(ppt_slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, 13.333, 7.5, bg)
        self._pptx_add_shape(ppt_slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0.0, 0.0, 0.1, 7.5, accent)
        self._pptx_add_textbox(ppt_slide, 0.72, 0.42, 4.9, 0.25, (slide.kicker or deck_title).upper()[:62], 9.5, color=muted, bold=True)
        self._pptx_add_textbox(ppt_slide, 12.05, 0.42, 0.7, 0.24, f"{index + 1} / {total}", 10, color=muted, bold=True, align=PP_ALIGN.RIGHT)
        self._pptx_add_textbox(ppt_slide, 0.72, 0.98, 11.5, 0.78, self._compact_title(slide.title, 8), 25, color=ink, bold=True)

        if slide.layout == "agenda":
            self._pptx_add_agenda_rows(ppt_slide, slide, 0.72, 2.02, 11.85, 4.55, direction)
        elif slide.layout == "metrics":
            self._pptx_add_metric_cards(ppt_slide, slide, 0.72, 2.03, 11.85, 4.42, direction)
        elif slide.layout == "timeline":
            self._pptx_add_timeline(ppt_slide, slide, 0.72, 2.18, 11.85, 3.95, direction)
        elif slide.layout == "compare" and slide.rows:
            self._pptx_add_table(ppt_slide, slide.rows, 0.72, 2.0, 11.85, 4.2, direction)
        elif slide.layout == "highlights":
            self._pptx_add_highlight_cards(ppt_slide, slide, 0.72, 2.03, 11.85, 4.42, direction)
        else:
            self._pptx_add_bullets(ppt_slide, (slide.bullets or slide.body or [])[:4], 0.9, 2.1, 11.2, 3.9, ink, accent)
        self._pptx_add_footer_line(ppt_slide, deck_title, direction)

    def _render_pptx_closing(self, ppt_slide: Any, slide: Slide, deck_title: str, index: int, total: int, folder: Path, direction: DeckVisualDirection) -> None:
        bg = self._pptx_color(direction.dark_bg)
        ink = self._pptx_color("#ffffff")
        muted = self._pptx_color(direction.dark_muted)
        accent = self._pptx_color(direction.accent)
        self._pptx_add_shape(ppt_slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, 13.333, 7.5, bg)
        self._pptx_add_shape(ppt_slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, 0.12, 7.5, accent)
        self._pptx_add_shape(ppt_slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0.72, 0.82, 0.86, 0.045, accent)
        self._pptx_add_textbox(ppt_slide, 0.72, 0.48, 4.8, 0.26, "MENSAGEM FINAL", 9.5, color=muted, bold=True)
        self._pptx_add_textbox(ppt_slide, 11.82, 0.48, 0.9, 0.24, f"{index + 1} / {total}", 10, color=muted, bold=True, align=PP_ALIGN.RIGHT)
        self._pptx_add_textbox(ppt_slide, 0.72, 1.48, 7.5, 1.25, self._compact_title(slide.title, 8), 31, color=ink, bold=True)
        items = (slide.bullets or slide.body or [])[:3]
        for idx, item in enumerate(items, start=1):
            y_pos = 3.2 + (idx - 1) * 0.74
            self._pptx_add_textbox(ppt_slide, 0.86, y_pos, 0.42, 0.22, f"{idx:02d}", 10, color=accent, bold=True)
            self._pptx_add_textbox(ppt_slide, 1.42, y_pos - 0.02, 6.7, 0.3, self._compact_copy(item, 13), 15, color=muted, bold=True)
        self._pptx_add_shape(ppt_slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 9.08, 1.38, 0.012, 4.42, self._pptx_color(direction.dark_line))
        self._pptx_add_textbox(ppt_slide, 9.44, 1.42, 2.9, 0.22, "PROXIMA ACAO", 9.5, color=accent, bold=True)
        self._pptx_add_textbox(ppt_slide, 9.44, 1.94, 2.8, 0.95, self._compact_copy(items[0] if items else slide.title, 14), 18, color=ink, bold=True)
        if slide.visual:
            self._pptx_try_add_image(ppt_slide, slide.visual, folder, 9.44, 4.72, 2.8, 1.02)

    def _pptx_add_visual_panel(self, ppt_slide: Any, slide: Slide, folder: Path, dark: bool, direction: DeckVisualDirection) -> None:
        image_added = False
        if slide.visual:
            image_added = self._pptx_try_add_image(ppt_slide, slide.visual, folder, 9.28, 1.02, 3.35, 4.35)
        if image_added:
            self._pptx_add_shape(
                ppt_slide,
                MSO_AUTO_SHAPE_TYPE.RECTANGLE,
                9.14,
                0.88,
                0.035,
                5.68,
                self._pptx_color(direction.accent),
            )
            self._pptx_add_textbox(ppt_slide, 9.28, 5.62, 3.35, 0.22, "VISUAL DE APOIO", 8.5, color=self._pptx_color(direction.muted), bold=True)
            return

        self._pptx_add_shape(
            ppt_slide,
            MSO_AUTO_SHAPE_TYPE.RECTANGLE,
            9.1,
            0.98,
            0.045,
            5.42,
            self._pptx_color(direction.accent),
        )
        self._pptx_add_textbox(ppt_slide, 9.42, 1.0, 2.8, 0.24, direction.focus_label, 8.5, color=self._pptx_color(direction.muted), bold=True)
        self._pptx_add_textbox(ppt_slide, 9.42, 1.48, 2.92, 0.82, self._compact_title(slide.title, 6), 17, color=self._pptx_color(direction.ink), bold=True)
        for idx, point in enumerate(self._pptx_visual_points(slide)[:3], start=1):
            y_pos = 2.72 + (idx - 1) * 0.86
            self._pptx_add_shape(ppt_slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 9.42, y_pos + 0.08, 0.2, 0.035, self._pptx_color(direction.accent))
            self._pptx_add_textbox(ppt_slide, 9.78, y_pos - 0.04, 2.65, 0.32, self._compact_copy(point, 8), 10.5, color=self._pptx_color(direction.muted), bold=True)
            self._pptx_add_shape(ppt_slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 9.42, y_pos + 0.5, 2.75, 0.01, self._pptx_color(direction.line))

    def _pptx_add_cover_art(self, ppt_slide: Any, slide: Slide, direction: DeckVisualDirection) -> None:
        self._pptx_add_shape(ppt_slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 9.08, 1.12, 0.035, 4.82, self._pptx_color(direction.accent))
        self._pptx_add_textbox(ppt_slide, 9.42, 1.12, 2.85, 0.22, direction.cover_label, 8.5, color=self._pptx_color(direction.dark_muted), bold=True)
        for idx, point in enumerate(self._pptx_visual_points(slide)[:3], start=1):
            y_pos = 1.72 + (idx - 1) * 1.08
            self._pptx_add_textbox(ppt_slide, 9.42, y_pos, 0.35, 0.18, f"{idx:02d}", 8.5, color=self._pptx_color(direction.accent), bold=True)
            self._pptx_add_textbox(ppt_slide, 9.98, y_pos - 0.02, 2.48, 0.38, self._compact_copy(point, 8), 11.5, color=self._pptx_color("#ffffff"), bold=True)
            self._pptx_add_shape(ppt_slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 9.42, y_pos + 0.58, 2.8, 0.01, self._pptx_color(direction.dark_line))

    def _pptx_try_add_image(self, ppt_slide: Any, visual: str, folder: Path, left: float, top: float, width: float, height: float) -> bool:
        try:
            if visual.startswith(("http://", "https://")):
                with httpx.Client(timeout=90, follow_redirects=True) as client:
                    response = client.get(visual)
                    response.raise_for_status()
                image_stream = BytesIO(response.content)
                ppt_slide.shapes.add_picture(image_stream, Inches(left), Inches(top), Inches(width), Inches(height))
                return True
            if visual.startswith("data:") and "," in visual:
                encoded = visual.split(",", 1)[1]
                ppt_slide.shapes.add_picture(BytesIO(base64.b64decode(encoded)), Inches(left), Inches(top), Inches(width), Inches(height))
                return True
            local_path = (folder / visual).resolve() if not Path(visual).is_absolute() else Path(visual)
            if local_path.exists():
                ppt_slide.shapes.add_picture(str(local_path), Inches(left), Inches(top), Inches(width), Inches(height))
                return True
        except Exception:
            return False
        return False

    def _pptx_add_agenda_rows(self, ppt_slide: Any, slide: Slide, left: float, top: float, width: float, height: float, direction: DeckVisualDirection) -> None:
        items = slide.bullets or slide.body or []
        item_count = min(len(items), 6)
        if not item_count:
            return
        row_height = min(0.72, height / item_count)
        self._pptx_add_shape(ppt_slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, left + 0.42, top + 0.12, 0.016, row_height * item_count, self._pptx_color(direction.line))
        for index, item in enumerate(items[:6], start=1):
            row_top = top + (index - 1) * row_height
            label, detail = self._split_agenda_item(item)
            self._pptx_add_textbox(ppt_slide, left, row_top + 0.04, 0.3, 0.16, f"{index:02d}", 9.5, color=self._pptx_color(direction.accent), bold=True, align=PP_ALIGN.RIGHT)
            self._pptx_add_shape(ppt_slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, left + 0.41, row_top + 0.16, 0.08, 0.08, self._pptx_color(direction.accent))
            self._pptx_add_textbox(ppt_slide, left + 0.72, row_top - 0.01, 3.25, 0.28, self._compact_title(label, 5), 15.5, color=self._pptx_color(direction.ink), bold=True)
            self._pptx_add_textbox(ppt_slide, left + 4.15, row_top + 0.02, width - 4.2, 0.24, self._compact_copy(detail, 12), 11.5, color=self._pptx_color(direction.muted))

    def _pptx_add_metric_cards(self, ppt_slide: Any, slide: Slide, left: float, top: float, width: float, height: float, direction: DeckVisualDirection) -> None:
        cards = (slide.bullets or slide.body or [])[:4]
        card_width = (width - 0.28) / 2
        card_height = (height - 0.26) / 2
        for index, item in enumerate(cards):
            col = index % 2
            row = index // 2
            card_left = left + col * (card_width + 0.28)
            card_top = top + row * (card_height + 0.26)
            self._pptx_add_shape(
                ppt_slide,
                MSO_AUTO_SHAPE_TYPE.RECTANGLE,
                card_left,
                card_top,
                card_width,
                card_height,
                self._pptx_color(direction.card_bg),
                line=direction.line,
            )
            self._pptx_add_shape(ppt_slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, card_left, card_top, card_width, 0.045, self._pptx_color(direction.accent))
            label, value, detail = self._parse_metric_bullet(f"- {item}")
            self._pptx_add_textbox(ppt_slide, card_left + 0.28, card_top + 0.26, card_width - 0.52, 0.22, label.upper()[:30], 9.5, color=self._pptx_color(direction.accent), bold=True)
            self._pptx_add_textbox(ppt_slide, card_left + 0.28, card_top + 0.62, card_width - 0.52, 0.44, value, 24, color=self._pptx_color(direction.ink), bold=True)
            self._pptx_add_textbox(ppt_slide, card_left + 0.28, card_top + 1.22, card_width - 0.52, 0.42, self._compact_copy(detail or label), 11, color=self._pptx_color(direction.muted))

    def _pptx_add_highlight_cards(self, ppt_slide: Any, slide: Slide, left: float, top: float, width: float, height: float, direction: DeckVisualDirection) -> None:
        cards = (slide.bullets or slide.body or [])[:4]
        for index, item in enumerate(cards):
            col = index % 2
            row = index // 2
            col_width = (width - 0.72) / 2
            block_left = left + col * (col_width + 0.72)
            block_top = top + row * 1.72
            label, detail = self._split_card_item(item)
            self._pptx_add_textbox(ppt_slide, block_left, block_top, 0.42, 0.18, f"{index + 1:02d}", 9.5, color=self._pptx_color(direction.accent), bold=True)
            self._pptx_add_shape(ppt_slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, block_left + 0.62, block_top + 0.06, 0.58, 0.035, self._pptx_color(direction.accent))
            self._pptx_add_textbox(ppt_slide, block_left + 0.62, block_top + 0.32, col_width - 0.66, 0.26, label.upper()[:34], 9.5, color=self._pptx_color(direction.accent), bold=True)
            self._pptx_add_textbox(ppt_slide, block_left + 0.62, block_top + 0.72, col_width - 0.66, 0.46, self._compact_title(detail or label, 7), 16, color=self._pptx_color(direction.ink), bold=True)
            self._pptx_add_textbox(ppt_slide, block_left + 0.62, block_top + 1.24, col_width - 0.66, 0.26, self._compact_copy(detail or label, 8), 9.5, color=self._pptx_color(direction.muted))
            self._pptx_add_shape(ppt_slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, block_left + 0.62, block_top + 1.6, col_width - 0.72, 0.012, self._pptx_color(direction.line))

    def _pptx_add_timeline(self, ppt_slide: Any, slide: Slide, left: float, top: float, width: float, height: float, direction: DeckVisualDirection) -> None:
        steps = (slide.bullets or slide.body or [])[:5]
        if not steps:
            return
        self._pptx_add_shape(ppt_slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, left + 0.34, top + 0.24, width - 0.68, 0.025, self._pptx_color(direction.line))
        step_width = (width - 0.34 * (len(steps) - 1)) / len(steps)
        for index, item in enumerate(steps, start=1):
            block_left = left + (index - 1) * (step_width + 0.34)
            self._pptx_add_shape(ppt_slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, block_left + 0.14, top + 0.11, 0.22, 0.22, self._pptx_color(direction.accent))
            self._pptx_add_textbox(ppt_slide, block_left + 0.13, top + 0.46, 0.34, 0.13, f"{index:02d}", 8.5, color=self._pptx_color(direction.accent), bold=True, align=PP_ALIGN.CENTER)
            self._pptx_add_shape(
                ppt_slide,
                MSO_AUTO_SHAPE_TYPE.RECTANGLE,
                block_left,
                top + 0.72,
                step_width,
                height - 0.92,
                self._pptx_color(direction.card_bg),
                line=direction.line,
            )
            self._pptx_add_textbox(ppt_slide, block_left + 0.22, top + 1.02, step_width - 0.44, 0.92, self._compact_title(item, 8), 13, color=self._pptx_color(direction.ink), bold=True)
            self._pptx_add_textbox(ppt_slide, block_left + 0.22, top + 2.05, step_width - 0.44, 0.44, self._compact_copy(item, 8), 9.5, color=self._pptx_color(direction.muted))

    def _pptx_add_table(self, ppt_slide: Any, rows: list[list[str]], left: float, top: float, width: float, height: float, direction: DeckVisualDirection) -> None:
        if not rows:
            return
        cols = max(len(row) for row in rows)
        table = ppt_slide.shapes.add_table(len(rows), cols, Inches(left), Inches(top), Inches(width), Inches(height)).table
        for row_index, row in enumerate(rows):
            for col_index in range(cols):
                cell = table.cell(row_index, col_index)
                text = row[col_index] if col_index < len(row) else ""
                cell.text = text
                cell.fill.solid()
                cell.fill.fore_color.rgb = self._pptx_color(direction.ink if row_index == 0 else direction.card_bg)
                cell.text_frame.word_wrap = True
                for paragraph in cell.text_frame.paragraphs:
                    paragraph.alignment = PP_ALIGN.LEFT
                    for run in paragraph.runs:
                        run.font.size = PptxPt(12 if row_index == 0 else 11)
                        run.font.bold = row_index == 0
                        run.font.color.rgb = self._pptx_color("#ffffff" if row_index == 0 else direction.ink)

    def _pptx_add_bullets(
        self,
        ppt_slide: Any,
        items: list[str],
        left: float,
        top: float,
        width: float,
        height: float,
        ink: PptxRGBColor,
        accent: PptxRGBColor,
    ) -> None:
        if not items:
            return
        for index, item in enumerate(items[:4]):
            y_pos = top + index * 0.88
            self._pptx_add_shape(ppt_slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, left, y_pos + 0.15, 0.16, 0.035, accent)
            self._pptx_add_textbox(ppt_slide, left + 0.36, y_pos, width - 0.36, min(height, 0.62), self._compact_copy(item, 16), 15, color=ink)

    def _pptx_add_chip(self, ppt_slide: Any, left: float, top: float, width: float, height: float, text: str, dark: bool) -> None:
        fill = "#34597d" if dark else "#e6eef6"
        font = "#ffffff" if dark else "#143458"
        self._pptx_add_shape(
            ppt_slide,
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
            left,
            top,
            width,
            height,
            self._pptx_color(fill),
            line=fill,
            transparency=0.1 if dark else 0.0,
        )
        self._pptx_add_textbox(
            ppt_slide,
            left + 0.08,
            top + 0.09,
            width - 0.16,
            height - 0.08,
            text,
            10,
            color=self._pptx_color(font),
            bold=True,
            align=PP_ALIGN.CENTER,
        )

    def _pptx_add_shape(
        self,
        ppt_slide: Any,
        shape_type: Any,
        left: float,
        top: float,
        width: float,
        height: float,
        fill_color: PptxRGBColor,
        line: str | None = None,
        transparency: float = 0.0,
    ) -> Any:
        shape = ppt_slide.shapes.add_shape(shape_type, Inches(left), Inches(top), Inches(width), Inches(height))
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
        if transparency:
            shape.fill.transparency = max(0.0, min(transparency, 1.0))
        if line is None:
            shape.line.fill.background()
        else:
            shape.line.color.rgb = self._pptx_color(line)
        return shape

    def _pptx_add_textbox(
        self,
        ppt_slide: Any,
        left: float,
        top: float,
        width: float,
        height: float,
        text: str,
        font_size: float,
        color: PptxRGBColor,
        bold: bool = False,
        align: PP_ALIGN = PP_ALIGN.LEFT,
    ) -> None:
        textbox = ppt_slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        frame = textbox.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.vertical_anchor = MSO_ANCHOR.TOP
        paragraph = frame.paragraphs[0]
        paragraph.alignment = align
        run = paragraph.add_run()
        run.text = text
        run.font.size = PptxPt(font_size)
        run.font.bold = bold
        run.font.name = "Aptos"
        run.font.color.rgb = color

    def _pptx_color(self, value: str) -> PptxRGBColor:
        normalized = value.strip().lstrip("#")
        return PptxRGBColor(int(normalized[0:2], 16), int(normalized[2:4], 16), int(normalized[4:6], 16))

    def _pptx_creative_direction(self, user_request: str, deck_title: str) -> DeckVisualDirection:
        text = self._normalize_for_match(f"{user_request} {deck_title}")
        directions = [
            (
                "food",
                ["comida", "alimento", "restaurante", "cardapio", "menu", "pizza", "hamburguer", "sushi", "cafe", "doce", "bolo", "gastronomia", "delivery", "cozinha"],
                DeckVisualDirection(
                    domain="food",
                    mood="sensorial, appetite-forward, warm editorial gastronomy",
                    dark_bg="#1a1009",
                    dark_divider="#3a2416",
                    light_bg="#fff8ef",
                    card_bg="#fffdf7",
                    ink="#2a160b",
                    muted="#7c6655",
                    accent="#e86f1f",
                    accent_2="#7a9f42",
                    line="#ead7c4",
                    dark_muted="#f0d6be",
                    dark_line="#4f3320",
                    cover_label="SABOR E EXPERIENCIA",
                    focus_label="NO PRATO",
                    image_style="premium food photography, natural light, appetizing composition, warm editorial restaurant aesthetic",
                ),
            ),
            (
                "animals",
                ["animal", "animais", "cachorro", "gato", "ave", "passaro", "peixe", "cavalo", "fauna", "selvagem", "veterinaria", "pet", "zoologia", "biologia", "habitat", "especie", "especies", "onca", "felino", "mamifero", "reptil", "inseto", "biodiversidade"],
                DeckVisualDirection(
                    domain="animals",
                    mood="natural, organic, field-guide clarity",
                    dark_bg="#0d1711",
                    dark_divider="#20382b",
                    light_bg="#f5faf2",
                    card_bg="#ffffff",
                    ink="#14251a",
                    muted="#627465",
                    accent="#4f9f52",
                    accent_2="#c08a2c",
                    line="#cfe1ce",
                    dark_muted="#c9dec9",
                    dark_line="#2c4a37",
                    cover_label="HABITAT E COMPORTAMENTO",
                    focus_label="EM CAMPO",
                    image_style="premium nature photography, habitat-aware, organic textures, documentary editorial style",
                ),
            ),
            (
                "fashion",
                ["moda", "fashion", "roupa", "look", "beleza", "estetica", "maquiagem", "marca pessoal", "luxo", "joia", "cosmetico"],
                DeckVisualDirection(
                    domain="fashion",
                    mood="elegant, editorial, high-fashion magazine",
                    dark_bg="#120d12",
                    dark_divider="#322332",
                    light_bg="#fbf7f9",
                    card_bg="#ffffff",
                    ink="#221821",
                    muted="#7a6877",
                    accent="#b65c86",
                    accent_2="#d8b36a",
                    line="#ead9e3",
                    dark_muted="#ead4df",
                    dark_line="#493143",
                    cover_label="ESTILO E DESEJO",
                    focus_label="NO LOOK",
                    image_style="premium fashion editorial, refined styling, magazine layout, elegant lighting, luxury visual language",
                ),
            ),
            (
                "health",
                ["saude", "saúde", "medico", "médico", "clinica", "clínica", "hospital", "bem-estar", "fitness", "nutricao", "nutrição", "terapia"],
                DeckVisualDirection(
                    domain="health",
                    mood="clean, trustworthy, calm clinical editorial",
                    dark_bg="#071718",
                    dark_divider="#14383b",
                    light_bg="#f3fbfa",
                    card_bg="#ffffff",
                    ink="#0f2b2d",
                    muted="#5f7778",
                    accent="#17a7a0",
                    accent_2="#75b86b",
                    line="#cbe6e3",
                    dark_muted="#c7e4e2",
                    dark_line="#244b4f",
                    cover_label="CUIDADO E CLAREZA",
                    focus_label="EM CUIDADO",
                    image_style="clean healthcare editorial, human-centered, calm light, credible medical wellness visual language",
                ),
            ),
            (
                "education",
                ["aula", "curso", "educacao", "educação", "escola", "ensino", "professor", "aprendizagem", "treinamento", "workshop", "didatico", "didático"],
                DeckVisualDirection(
                    domain="education",
                    mood="clear, didactic, modern learning system",
                    dark_bg="#0d1426",
                    dark_divider="#25345e",
                    light_bg="#f6f8ff",
                    card_bg="#ffffff",
                    ink="#14213d",
                    muted="#65718f",
                    accent="#4969e8",
                    accent_2="#f0b429",
                    line="#d9e0fb",
                    dark_muted="#ccd6ff",
                    dark_line="#31427a",
                    cover_label="TRILHA DE APRENDIZADO",
                    focus_label="EM AULA",
                    image_style="modern education editorial, clear diagrams, calm study environment, premium learning design",
                ),
            ),
            (
                "technology",
                ["tecnologia", "software", "ia", "inteligencia artificial", "dados", "api", "sistema", "app", "cloud", "cyber", "seguranca", "segurança", "automacao"],
                DeckVisualDirection(
                    domain="technology",
                    mood="precise, modern, systems-oriented",
                    dark_bg="#07131b",
                    dark_divider="#173344",
                    light_bg="#f3f8fa",
                    card_bg="#ffffff",
                    ink="#102631",
                    muted="#607681",
                    accent="#18a4a6",
                    accent_2="#62c370",
                    line="#cde5e5",
                    dark_muted="#c0dde0",
                    dark_line="#254a59",
                    cover_label="SISTEMA E DECISAO",
                    focus_label="EM SISTEMA",
                    image_style="premium technology editorial, precise interface details, subtle depth, modern systems visual language",
                ),
            ),
            (
                "business",
                ["empresa", "negocio", "negócio", "vendas", "marketing", "financeiro", "gestao", "gestão", "estrategia", "estratégia", "produto", "startup", "mercado"],
                DeckVisualDirection(
                    domain="business",
                    mood="executive, strategic, consulting-grade",
                    dark_bg="#06120f",
                    dark_divider="#17352d",
                    light_bg="#f7faf4",
                    card_bg="#ffffff",
                    ink="#0d241c",
                    muted="#667a70",
                    accent="#15966f",
                    accent_2="#d6a33d",
                    line="#d7e7df",
                    dark_muted="#b7cec5",
                    dark_line="#23453d",
                    cover_label="LEITURA DO DECK",
                    focus_label="EM FOCO",
                    image_style="premium consulting presentation visual language, editorial business photography, restrained executive composition",
                ),
            ),
        ]
        for _, keywords, direction in directions:
            if any(keyword in text for keyword in keywords):
                return direction
        return DeckVisualDirection(
            domain="adaptive",
            mood="adaptive, editorial, user-intent aware",
            dark_bg="#101114",
            dark_divider="#2d3035",
            light_bg="#f8f8f4",
            card_bg="#ffffff",
            ink="#202124",
            muted="#6c706d",
            accent="#5d8f6a",
            accent_2="#b98e42",
            line="#deded6",
            dark_muted="#d8d8ce",
            dark_line="#3c403d",
            cover_label="IDEIA CENTRAL",
            focus_label="EM CONTEXTO",
            image_style="premium editorial presentation visual language, subject-specific, tasteful, minimal, no generic template look",
        )

    def _normalize_for_match(self, value: str) -> str:
        normalized = value.lower()
        replacements = {
            "á": "a",
            "à": "a",
            "â": "a",
            "ã": "a",
            "é": "e",
            "ê": "e",
            "í": "i",
            "ó": "o",
            "ô": "o",
            "õ": "o",
            "ú": "u",
            "ç": "c",
        }
        for old, new in replacements.items():
            normalized = normalized.replace(old, new)
        return normalized

    def _truncate_words(self, text: str, count: int) -> str:
        words = self._clean_inline_markdown(text).split()
        return "\n".join(words[:count]) if words else "Tema"

    def _compact_title(self, text: str, max_words: int = 9) -> str:
        cleaned = self._clean_inline_markdown(text).strip()
        words = cleaned.split()
        if len(words) <= max_words:
            return cleaned or "Apresentacao"
        return " ".join(words[:max_words]).rstrip(".,;:") + "."

    def _compact_copy(self, text: str, max_words: int = 8) -> str:
        words = self._clean_inline_markdown(text).split()
        compact = " ".join(words[:max_words]).strip()
        if len(words) > max_words:
            compact += "..."
        return compact or text

    def _pptx_add_footer_line(self, ppt_slide: Any, deck_title: str, direction: DeckVisualDirection) -> None:
        self._pptx_add_shape(ppt_slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0.72, 6.92, 0.6, 0.04, self._pptx_color(direction.accent))
        self._pptx_add_textbox(ppt_slide, 1.44, 6.83, 5.6, 0.22, self._compact_copy(deck_title, 8).upper(), 8, color=self._pptx_color(direction.muted), bold=True)

    def _pptx_visual_heading(self, slide: Slide) -> str:
        for source in [slide.body or [], slide.bullets or []]:
            for item in source:
                text = self._clean_inline_markdown(item).strip()
                if text:
                    return self._truncate_words(text, 6)
        return self._truncate_words(slide.title, 6)

    def _pptx_visual_points(self, slide: Slide) -> list[str]:
        points: list[str] = []
        source = slide.bullets or slide.body or []
        for item in source:
            label, detail = self._split_card_item(item)
            candidate = detail if detail and detail != "Ponto principal do slide" else label
            candidate = self._clean_inline_markdown(candidate).strip()
            if candidate:
                points.append(self._truncate_words(candidate, 7).replace("\n", " "))
        if not points:
            points.append(self._truncate_words(slide.title, 7).replace("\n", " "))
        while len(points) < 3:
            points.append("Direcao clara para decisao executiva")
        return points[:3]

    def _humanize_slide_title(self, raw_title: str, index: int, deck_title: str) -> str:
        cleaned = self._clean_inline_markdown(re.sub(r"^#{1,3}\s*", "", raw_title or "")).strip()
        cleaned = re.sub(r"^(slide|pagina|página|page)\s*\d+\s*[-–:]\s*", "", cleaned, flags=re.I)
        cleaned = re.sub(r"^\d+\s*[-–.:]\s*", "", cleaned)
        return cleaned or (deck_title if index == 0 else f"Topico {index + 1}")

    def _extract_slide_table(self, raw_slide: str) -> list[list[str]] | None:
        lines = [line.strip() for line in raw_slide.splitlines() if line.strip()]
        table_lines = [line for line in lines if self._looks_like_table_line(line)]
        rows = self._parse_markdown_table(table_lines) if table_lines else []
        return rows or None

    def _slide_kicker(self, deck_title: str, layout: str, index: int) -> str:
        mapping = {
            "lead": deck_title,
            "agenda": "Visao do percurso",
            "metrics": "Sinais que importam",
            "timeline": "Leitura em etapas",
            "compare": "Comparacao executiva",
            "highlights": "Prioridades executivas",
            "closing": "Mensagem final",
            "section": "Ponto de analise",
        }
        return mapping.get(layout, f"Slide {index + 1}")

    def _slide_visual_fallback_url(self, deck_title: str, raw_slide: str, index: int, total: int) -> str | None:
        return None

    def _build_presentation_html(self, deck_title: str, slides: list[Slide], user_request: str = "") -> str:
        theme_css = self._deck_theme_css(user_request, deck_title)
        slide_markup = "\n".join(self._render_slide_html(slide, index, len(slides)) for index, slide in enumerate(slides))
        return f"""<!doctype html>
<html lang="pt-BR">
<head>
  <meta charset="utf-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1" />
  <title>{self._escape_html(deck_title)}</title>
  <style>
    :root {{
      color-scheme: light;
      --bg: #0f1d2e;
      --panel: rgba(255,255,255,0.88);
      --panel-soft: rgba(255,255,255,0.7);
      --line: rgba(17, 35, 58, 0.09);
      --ink: #10233d;
      --muted: #61748f;
      --accent: #1f5ea8;
      --accent-2: #27a6b8;
      --shadow: 0 28px 60px rgba(14, 31, 53, 0.12);
      --radius: 28px;
      --slide-w: 1920px;
      --slide-h: 1080px;
    }}
    {theme_css}
    @page {{
      size: 1920px 1080px;
      margin: 0;
    }}
    * {{
      box-sizing: border-box;
    }}
    html, body {{
      width: 1920px;
      margin: 0;
      padding: 0;
      background: #0b1320;
      color: var(--ink);
      font-family: Inter, Aptos, "Segoe UI", Arial, sans-serif;
    }}
    body {{
      display: block;
      padding: 24px;
    }}
    .deck {{
      width: var(--slide-w);
      display: grid;
      gap: 26px;
    }}
    body.pdf-export {{
      width: 1920px !important;
      height: 1080px !important;
      display: block;
      padding: 0 !important;
      margin: 0 !important;
      overflow: hidden !important;
      background: #fff;
    }}
    body.pdf-export .deck {{
      width: 1920px !important;
      gap: 0;
    }}
    .slide-container {{
      width: var(--slide-w);
      height: var(--slide-h);
      position: relative;
      overflow: hidden;
      break-inside: avoid;
      page-break-after: always;
    }}
    .slide-container:last-child {{
      page-break-after: auto;
    }}
    body.pdf-export .slide {{
      width: 1920px !important;
      height: 1080px !important;
      aspect-ratio: auto;
      border-radius: 0;
      box-shadow: none;
      page-break-after: always;
      break-after: page;
    }}
    body.pdf-export .slide:last-child {{
      page-break-after: auto;
      break-after: auto;
    }}
    body.pdf-export .deck-nav {{
      display: none !important;
    }}
    .slide {{
      position: relative;
      width: var(--slide-w);
      height: var(--slide-h);
      overflow: hidden;
      border-radius: 30px;
      background:
        radial-gradient(circle at top left, rgba(31, 94, 168, 0.22), transparent 34%),
        linear-gradient(180deg, #f7fbff 0%, #eef4fa 100%);
      box-shadow: 0 42px 90px rgba(8, 19, 36, 0.28);
      display: grid;
      grid-template-columns: minmax(0, 1.08fr) minmax(320px, .92fr);
      isolation: isolate;
    }}
    .slide-content {{
      padding: 72px 84px 72px;
      display: flex;
      flex-direction: column;
      min-width: 0;
      gap: 24px;
      z-index: 2;
    }}
    .slide-meta {{
      display: flex;
      align-items: center;
      justify-content: space-between;
      color: var(--muted);
      font-size: 15px;
      font-weight: 600;
      letter-spacing: .01em;
    }}
    .slide-kicker {{
      text-transform: uppercase;
      letter-spacing: .12em;
      font-size: 18px;
      color: var(--accent);
      font-weight: 800;
    }}
    .slide-page {{
      font-size: 20px;
    }}
    .slide-title {{
      margin: 0;
      font-size: 76px;
      line-height: 1.02;
      letter-spacing: 0;
      color: #0c2848;
      max-width: 11ch;
      text-wrap: balance;
    }}
    .slide-subtitle {{
      margin: 0;
      font-size: 32px;
      line-height: 1.35;
      color: #41556f;
      max-width: 27ch;
      text-wrap: balance;
    }}
    .slide-visual {{
      position: relative;
      min-width: 0;
      background: linear-gradient(135deg, #123255 0%, #1a5d85 100%);
    }}
    .slide-visual img {{
      width: 100%;
      height: 100%;
      object-fit: cover;
      display: block;
    }}
    .slide-visual::after {{
      content: "";
      position: absolute;
      inset: 0;
      background: linear-gradient(90deg, rgba(247,251,255,0.02) 0%, rgba(12, 28, 49, 0.18) 100%);
    }}
    .visual-map {{
      position: absolute;
      inset: 54px;
      display: grid;
      place-items: center;
      color: rgba(255,255,255,0.92);
    }}
    .visual-orbit {{
      position: relative;
      width: 360px;
      height: 360px;
      border-radius: 999px;
      border: 1px solid rgba(255,255,255,0.22);
      background:
        radial-gradient(circle at center, rgba(255,255,255,0.20), transparent 28%),
        radial-gradient(circle at 30% 20%, rgba(94,234,212,0.26), transparent 20%);
    }}
    .visual-node {{
      position: absolute;
      width: 118px;
      min-height: 54px;
      border-radius: 999px;
      display: grid;
      place-items: center;
      text-align: center;
      padding: 10px 14px;
      background: rgba(255,255,255,0.14);
      border: 1px solid rgba(255,255,255,0.20);
      font-size: 13px;
      font-weight: 800;
      letter-spacing: .04em;
      text-transform: uppercase;
      backdrop-filter: blur(14px);
    }}
    .visual-node:nth-child(1) {{
      top: -8px;
      left: 50%;
      transform: translateX(-50%);
    }}
    .visual-node:nth-child(2) {{
      right: -38px;
      top: 48%;
      transform: translateY(-50%);
    }}
    .visual-node:nth-child(3) {{
      bottom: -8px;
      left: 50%;
      transform: translateX(-50%);
    }}
    .visual-node:nth-child(4) {{
      left: -38px;
      top: 48%;
      transform: translateY(-50%);
    }}
    .visual-core {{
      position: absolute;
      inset: 122px;
      border-radius: 999px;
      display: grid;
      place-items: center;
      text-align: center;
      padding: 20px;
      background: rgba(255,255,255,0.92);
      color: #0c2848;
      font-size: 20px;
      font-weight: 900;
      box-shadow: 0 22px 54px rgba(4, 17, 34, 0.20);
    }}
    .content-stack {{
      display: flex;
      flex-direction: column;
      gap: 14px;
      min-width: 0;
    }}
    .bullet-list {{
      list-style: none;
      margin: 0;
      padding: 0;
      display: grid;
      gap: 12px;
    }}
    .bullet-list li {{
      display: grid;
      grid-template-columns: 20px minmax(0, 1fr);
      gap: 20px;
      align-items: start;
      font-size: 32px;
      line-height: 1.36;
      color: #132942;
    }}
    .bullet-list li::before {{
      content: "";
      width: 20px;
      height: 20px;
      margin-top: 14px;
      border-radius: 999px;
      background: linear-gradient(135deg, var(--accent), var(--accent-2));
      box-shadow: 0 0 0 8px rgba(31, 94, 168, 0.08);
    }}
    .agenda-grid, .metrics-grid, .highlights-grid {{
      display: grid;
      grid-template-columns: repeat(2, minmax(0, 1fr));
      gap: 24px;
    }}
    .agenda-list {{
      display: grid;
      gap: 14px;
      align-content: start;
    }}
    .agenda-row {{
      display: grid;
      grid-template-columns: 58px minmax(0, 1fr);
      gap: 18px;
      align-items: start;
      padding: 20px 22px;
      border-radius: 24px;
      background: rgba(255,255,255,0.72);
      border: 1px solid rgba(16, 35, 61, 0.08);
      box-shadow: 0 10px 26px rgba(12, 40, 72, 0.08);
    }}
    .agenda-index {{
      display: grid;
      place-items: center;
      width: 58px;
      height: 58px;
      border-radius: 999px;
      background: linear-gradient(135deg, var(--accent), var(--accent-2));
      color: #fff;
      font-size: 18px;
      font-weight: 900;
      letter-spacing: .08em;
    }}
    .agenda-copy {{
      display: grid;
      gap: 8px;
    }}
    .agenda-copy strong {{
      font-size: 18px;
      text-transform: uppercase;
      letter-spacing: .12em;
      color: var(--accent);
    }}
    .agenda-copy span {{
      display: block;
      font-size: 34px;
      line-height: 1.18;
      color: #10233d;
      font-weight: 700;
      text-wrap: balance;
    }}
    .agenda-card, .metric-card, .highlight-card {{
      background: var(--panel);
      border: 1px solid var(--line);
      border-radius: 20px;
      padding: 20px 22px;
      box-shadow: var(--shadow);
      min-height: 0;
    }}
    .highlight-card {{
      background: linear-gradient(180deg, rgba(255,255,255,0.96), rgba(244,249,255,0.94));
      min-height: 150px;
      display: flex;
      flex-direction: column;
      justify-content: space-between;
      gap: 12px;
    }}
    .agenda-card strong, .metric-label {{
      display: block;
      font-size: 15px;
      text-transform: uppercase;
      letter-spacing: .08em;
      color: var(--accent);
      margin-bottom: 12px;
    }}
    .agenda-card span {{
      display: block;
      font-size: 24px;
      line-height: 1.22;
      color: #10233d;
      font-weight: 700;
      text-wrap: balance;
    }}
    .metric-value {{
      display: block;
      font-size: 34px;
      line-height: 1;
      letter-spacing: -.04em;
      color: #0c2848;
      margin: 8px 0 10px;
      font-weight: 800;
    }}
    .metric-text {{
      font-size: 18px;
      line-height: 1.45;
      color: #4b5e78;
    }}
    .highlight-label {{
      display: inline-flex;
      width: fit-content;
      padding: 8px 12px;
      border-radius: 999px;
      background: rgba(31, 94, 168, 0.08);
      color: var(--accent);
      font-size: 12px;
      font-weight: 800;
      letter-spacing: .08em;
      text-transform: uppercase;
    }}
    .highlight-text {{
      font-size: 22px;
      line-height: 1.24;
      color: #10233d;
      font-weight: 700;
      letter-spacing: -.03em;
    }}
    .highlight-detail {{
      font-size: 15px;
      line-height: 1.45;
      color: #51657f;
    }}
    .timeline {{
      display: grid;
      gap: 16px;
    }}
    .timeline-step {{
      position: relative;
      padding: 0 0 0 68px;
      min-height: 52px;
      display: flex;
      align-items: center;
      font-size: 21px;
      line-height: 1.34;
      color: #10233d;
    }}
    .timeline-step::before {{
      content: attr(data-step);
      position: absolute;
      left: 0;
      top: 0;
      width: 46px;
      height: 46px;
      border-radius: 999px;
      background: linear-gradient(135deg, #123d72, #1da6b7);
      color: #fff;
      display: grid;
      place-items: center;
      font-weight: 800;
      box-shadow: 0 12px 24px rgba(16, 42, 72, 0.18);
    }}
    .timeline-step::after {{
      content: "";
      position: absolute;
      left: 22px;
      top: 46px;
      width: 2px;
      bottom: -18px;
      background: rgba(18, 61, 114, 0.18);
    }}
    .timeline-step:last-child::after {{
      display: none;
    }}
    .compare-table {{
      width: 100%;
      border-collapse: collapse;
      overflow: hidden;
      border-radius: 24px;
      background: rgba(255,255,255,0.92);
      box-shadow: var(--shadow);
    }}
    .compare-table th, .compare-table td {{
      padding: 16px 18px;
      text-align: left;
      vertical-align: top;
      border-bottom: 1px solid rgba(16, 35, 61, 0.08);
      font-size: 18px;
      line-height: 1.4;
    }}
    .compare-table th {{
      background: #10345e;
      color: #fff;
      font-size: 15px;
      letter-spacing: .04em;
      text-transform: uppercase;
    }}
    .compare-table tr:last-child td {{
      border-bottom: 0;
    }}
    .lead {{
      grid-template-columns: minmax(0, 1.02fr) minmax(320px, .98fr);
      background:
        radial-gradient(circle at top left, rgba(36, 172, 196, 0.24), transparent 26%),
        linear-gradient(135deg, #0e1f33 0%, #163658 52%, #215b83 100%);
    }}
    .lead .slide-content {{
      padding-top: 84px;
      padding-bottom: 84px;
      justify-content: center;
    }}
    .lead .slide-kicker,
    .lead .slide-page,
    .lead .slide-subtitle {{
      color: rgba(255,255,255,0.8);
    }}
    .lead .slide-page {{
      color: rgba(255,255,255,0.9);
    }}
    .lead .slide-title {{
      color: #fff;
      font-size: 88px;
      max-width: 8.8ch;
    }}
    .lead .slide-subtitle {{
      font-size: 34px;
      max-width: 22ch;
    }}
    .lead .lead-chip {{
      display: inline-flex;
      align-items: center;
      gap: 12px;
      padding: 16px 20px;
      border-radius: 999px;
      background: rgba(255,255,255,0.08);
      color: rgba(255,255,255,0.92);
      font-size: 20px;
      width: fit-content;
      backdrop-filter: blur(12px);
    }}
    .lead-points {{
      display: flex;
      flex-wrap: wrap;
      gap: 14px;
      margin-top: 4px;
    }}
    .lead-point {{
      display: inline-flex;
      align-items: center;
      padding: 14px 18px;
      border-radius: 999px;
      background: rgba(255,255,255,0.08);
      border: 1px solid rgba(255,255,255,0.12);
      color: rgba(255,255,255,0.92);
      font-size: 20px;
      font-weight: 700;
      line-height: 1.2;
      max-width: 420px;
      text-wrap: balance;
    }}
    .lead .bullet-list li {{
      color: rgba(255,255,255,0.95);
      font-size: 24px;
      line-height: 1.3;
    }}
    .lead .bullet-list li::before {{
      background: linear-gradient(135deg, #68ecff, #d8faff);
      box-shadow: 0 0 0 8px rgba(255,255,255,0.08);
    }}
    .closing {{
      background:
        radial-gradient(circle at top right, rgba(39, 166, 184, 0.18), transparent 28%),
        linear-gradient(135deg, #10243d 0%, #1b4977 100%);
    }}
    .closing .slide-title,
    .closing .slide-subtitle,
    .closing .slide-kicker,
    .closing .slide-page,
    .closing .bullet-list li {{
      color: #fff;
    }}
    .closing .bullet-list li::before {{
      background: linear-gradient(135deg, #4fe0f2, #ffffff);
      box-shadow: 0 0 0 8px rgba(255,255,255,0.08);
    }}
    .content {{
      grid-template-columns: minmax(0, 1.18fr) minmax(320px, .82fr);
    }}
    .content .slide-title {{
      max-width: 14ch;
    }}
    .content.no-visual,
    .closing.no-visual,
    .lead.no-visual {{
      grid-template-columns: 1fr;
    }}
    .content.no-visual .slide-content,
    .closing.no-visual .slide-content {{
      max-width: 1440px;
    }}
    .slide.no-visual .slide-visual {{
      display: none;
    }}
    .slide.copy-heavy .bullet-list li {{
      font-size: 28px;
    }}
    .slide.copy-heavy .slide-title {{
      font-size: 68px;
    }}
    .slide.long-title .slide-title {{
      font-size: 66px;
      max-width: 18ch;
    }}
    .lead.long-title .slide-title {{
      font-size: 80px;
      max-width: 13ch;
    }}
    .deck-nav {{
      position: fixed;
      right: 24px;
      bottom: 22px;
      z-index: 20;
      display: flex;
      gap: 10px;
      background: rgba(9, 18, 32, 0.84);
      color: #fff;
      padding: 10px 14px;
      border-radius: 999px;
      box-shadow: 0 16px 30px rgba(0,0,0,0.24);
      backdrop-filter: blur(14px);
      font-size: 14px;
    }}
    .deck-nav button {{
      border: 0;
      background: rgba(255,255,255,0.08);
      color: inherit;
      padding: 10px 12px;
      border-radius: 999px;
      cursor: pointer;
      font: inherit;
    }}
    .deck-nav button:hover {{
      background: rgba(255,255,255,0.16);
    }}
    @media screen and (max-width: 1100px) {{
      body {{
        padding: 10px;
      }}
      .deck {{
        gap: 12px;
      }}
      .slide-container,
      .slide {{
        width: 1920px;
        height: 1080px;
        border-radius: 18px;
      }}
    }}
    @media print {{
      @page {{
        size: 1920px 1080px;
        margin: 0;
      }}
      * {{
        -webkit-print-color-adjust: exact !important;
        print-color-adjust: exact !important;
      }}
      html,
      body {{
        width: 1920px !important;
        height: 1080px !important;
        background: #fff;
        padding: 0;
        margin: 0;
        overflow: hidden;
      }}
      .deck {{
        width: 1920px !important;
        gap: 0;
      }}
      .slide-container {{
        width: 1920px !important;
        height: 1080px !important;
        overflow: hidden;
        page-break-after: always;
        break-after: page;
        break-inside: avoid;
      }}
      .slide {{
        width: 1920px !important;
        height: 1080px !important;
        margin: 0;
        border-radius: 0;
        box-shadow: none;
      }}
      .deck-nav {{
        display: none !important;
      }}
    }}
  </style>
</head>
<body>
  <main class="deck">
    {slide_markup}
  </main>
  <div class="deck-nav" aria-hidden="true">
    <button onclick="stepSlide(-1)">Anterior</button>
    <button onclick="stepSlide(1)">Próximo</button>
  </div>
  <script>
    const slides = Array.from(document.querySelectorAll('.slide'));
    let activeSlide = 0;
    function stepSlide(direction) {{
      activeSlide = Math.max(0, Math.min(slides.length - 1, activeSlide + direction));
      slides[activeSlide].scrollIntoView({{ behavior: 'smooth', block: 'center' }});
    }}
    document.addEventListener('keydown', (event) => {{
      if (event.key === 'ArrowRight' || event.key === 'PageDown') stepSlide(1);
      if (event.key === 'ArrowLeft' || event.key === 'PageUp') stepSlide(-1);
    }});
  </script>
</body>
</html>"""

    def _render_slide_html(self, slide: Slide, index: int, total: int) -> str:
        classes = [slide.layout if slide.layout in {"lead", "agenda", "metrics", "timeline", "compare", "closing", "highlights"} else "content"]
        if not slide.visual:
            if slide.layout in {"agenda", "metrics", "timeline", "compare", "highlights"}:
                classes.append("no-visual")
            else:
                classes.append("generated-visual")
        if len(slide.title) > 44:
            classes.append("long-title")
        if len(slide.bullets or []) >= 4:
            classes.append("copy-heavy")
        content = self._render_slide_content(slide)
        visual_html = self._render_visual_html(slide)
        title = self._escape_html(slide.title)
        kicker = self._escape_html(slide.kicker)
        subtitle = self._escape_html(slide.body[0]) if slide.layout == "lead" and slide.body else ""
        lead_chip = '<div class="lead-chip">Deck personalizado para leitura executiva</div>' if slide.layout == "lead" else ""
        return (
            f'<article class="slide-container"><section class="slide {" ".join(classes)}" id="slide-{index + 1}">'
            '<div class="slide-content">'
            f'<div class="slide-meta"><span class="slide-kicker" contenteditable="true">{kicker}</span><span class="slide-page">{index + 1} / {total}</span></div>'
            f'<h1 class="slide-title" contenteditable="true">{title}</h1>'
            + (f'<p class="slide-subtitle" contenteditable="true">{subtitle}</p>' if subtitle else "")
            + lead_chip
            + f'<div class="content-stack">{content}</div>'
            + '</div>'
            + visual_html
            + '</section></article>'
        )

    def _render_visual_html(self, slide: Slide) -> str:
        if not slide.visual:
            return self._render_generated_visual_html(slide)
        if slide.visual.startswith(("http://", "https://", "data:")):
            visual = self._escape_html(slide.visual)
        else:
            visual = quote(slide.visual.replace("\\", "/"), safe="/:.-_")
        return f'<aside class="slide-visual"><img src="{visual}" alt="{self._escape_html(slide.title)}" /></aside>'

    def _render_generated_visual_html(self, slide: Slide) -> str:
        """Render a clean, premium icon-grid visual when no image is available."""
        items = (slide.bullets or slide.body or [])[:4]
        icons = ["◈", "◉", "◆", "◇"]
        cards_html = ""
        for idx, item in enumerate(items):
            label = item[:32]
            icon = icons[idx % len(icons)]
            cards_html += f'<div class="vg-card"><span class="vg-icon">{icon}</span><span class="vg-label">{label}</span></div>'
        return f'<aside class="slide-visual"><div class="visual-grid">{cards_html}</div></aside>'

    def _render_slide_content(self, slide: Slide) -> str:
        if slide.layout == "lead":
            items = (slide.bullets or [])[:2]
            if not items:
                return ""
            chips = "".join(f'<span class="lead-point" contenteditable="true">{self._escape_html(item)}</span>' for item in items)
            return f'<div class="lead-points">{chips}</div>'
        if slide.layout == "agenda":
            items = slide.bullets or slide.body or []
            rows = []
            for idx, item in enumerate(items[:5], start=1):
                label, detail = self._split_agenda_item(item)
                rows.append(
                    '<article class="agenda-row">'
                    f'<span class="agenda-index">{idx:02d}</span>'
                    '<div class="agenda-copy">'
                    f'<strong contenteditable="true">{self._escape_html(label)}</strong>'
                    f'<span contenteditable="true">{self._escape_html(detail)}</span>'
                    '</div>'
                    '</article>'
                )
            return f'<div class="agenda-list">{"".join(rows)}</div>'
        if slide.layout == "metrics":
            cards = []
            for bullet in (slide.bullets or [])[:4]:
                label, value, detail = self._parse_metric_bullet(f"- {bullet}")
                cards.append(
                    '<article class="metric-card">'
                    f'<span class="metric-label" contenteditable="true">{self._escape_html(label)}</span>'
                    f'<span class="metric-value" contenteditable="true">{self._escape_html(value)}</span>'
                    f'<span class="metric-text" contenteditable="true">{self._escape_html(detail)}</span>'
                    '</article>'
                )
            return f'<div class="metrics-grid">{"".join(cards)}</div>'
        if slide.layout == "timeline":
            steps = []
            for idx, bullet in enumerate((slide.bullets or [])[:5], start=1):
                steps.append(f'<div class="timeline-step" data-step="{idx}" contenteditable="true">{self._escape_html(bullet)}</div>')
            return f'<div class="timeline">{"".join(steps)}</div>'
        if slide.layout == "highlights":
            cards = []
            for bullet in (slide.bullets or [])[:4]:
                label, detail = self._split_card_item(bullet)
                cards.append(
                    '<article class="highlight-card">'
                    f'<span class="highlight-label">{self._escape_html(label)}</span>'
                    f'<div class="highlight-text">{self._escape_html(detail)}</div>'
                    f'<div class="highlight-detail">{self._escape_html(self._short_support_copy(detail))}</div>'
                    '</article>'
                )
            return f'<div class="highlights-grid">{"".join(cards)}</div>'
        if slide.layout == "compare" and slide.rows:
            return self._render_custom_table(slide.rows)
        items = slide.bullets or slide.body or []
        if items:
            bullets = "".join(f'<li contenteditable="true"><span>{self._format_inline_html(item)}</span></li>' for item in items)
            return f'<ul class="bullet-list">{bullets}</ul>'
        return f'<p class="slide-subtitle" contenteditable="true">{self._escape_html(slide.kicker)}</p>'

    def _render_custom_table(self, rows: list[list[str]]) -> str:
        header = "".join(f'<th contenteditable="true">{self._escape_html(cell)}</th>' for cell in rows[0])
        body = "".join(
            "<tr>" + "".join(f'<td contenteditable="true">{self._format_inline_html(cell)}</td>' for cell in row) + "</tr>"
            for row in rows[1:]
        )
        return f'<table class="compare-table"><thead><tr>{header}</tr></thead><tbody>{body}</tbody></table>'

    def _split_agenda_item(self, item: str) -> tuple[str, str]:
        cleaned = self._clean_inline_markdown(item)
        if " - " in cleaned:
            left, right = cleaned.split(" - ", 1)
            return left.strip(), right.strip()
        words = cleaned.split()
        if len(words) <= 2:
            return cleaned, "Ponto central do percurso"
        return words[0], " ".join(words[1:])

    def _split_card_item(self, item: str) -> tuple[str, str]:
        cleaned = self._clean_inline_markdown(item)
        for token in [" - ", ": ", " — ", " -- "]:
            if token in cleaned:
                left, right = cleaned.split(token, 1)
                return left.strip()[:48], right.strip()
        words = cleaned.split()
        if len(words) <= 3:
            return cleaned, "Ponto principal do slide"
        return " ".join(words[:2]), " ".join(words[2:])

    def _deck_theme_css(self, user_request: str, deck_title: str) -> str:
        lowered = f"{user_request}\n{deck_title}".lower()
        themes = {
            "clear-green": """
    :root {
      --bg: #eef7f1;
      --panel: rgba(255,255,255,0.96);
      --panel-soft: rgba(255,255,255,0.82);
      --line: rgba(15, 23, 42, 0.08);
      --ink: #0f172a;
      --muted: #5b6472;
      --accent: #12715b;
      --accent-2: #38b27a;
      --shadow: 0 24px 52px rgba(15, 23, 42, 0.10);
    }
            """,
            "neon-blue": """
    :root {
      --bg: #050816;
      --panel: rgba(10,18,34,0.74);
      --panel-soft: rgba(15,24,44,0.58);
      --line: rgba(148,163,184,0.12);
      --ink: #e5eefb;
      --muted: #a8bad5;
      --accent: #60a5fa;
      --accent-2: #22d3ee;
      --shadow: 0 28px 60px rgba(2, 6, 23, 0.38);
    }
    .slide {
      background:
        radial-gradient(circle at top left, rgba(96, 165, 250, 0.22), transparent 34%),
        linear-gradient(180deg, #08101f 0%, #0d1a31 100%);
    }
    .slide-title, .agenda-card span, .highlight-text, .timeline-step, .bullet-list li, .metric-value {
      color: #f8fbff;
    }
    .slide-subtitle, .metric-text, .highlight-detail {
      color: #c7d6ea;
    }
    .agenda-card, .metric-card, .highlight-card, .compare-table {
      background: rgba(8,16,31,0.66);
    }
            """,
            "executive-dark": """
    :root {
      --bg: #0f1d2e;
      --panel: rgba(255,255,255,0.88);
      --panel-soft: rgba(255,255,255,0.7);
      --line: rgba(17, 35, 58, 0.09);
      --ink: #10233d;
      --muted: #61748f;
      --accent: #1f5ea8;
      --accent-2: #27a6b8;
      --shadow: 0 28px 60px rgba(14, 31, 53, 0.12);
    }
            """,
        }
        if any(token in lowered for token in ["neon", "futurista", "blue", "azul", "cyber"]):
            return themes["neon-blue"]
        if any(token in lowered for token in ["green", "verde", "claro", "clean", "minimal", "minimalista"]):
            return themes["clear-green"]
        return themes["executive-dark"]

    def _inline_slide_assets(self, html: str, base_dir: Path) -> str:
        def replace_src(match: re.Match[str]) -> str:
            original = match.group(1)
            if original.startswith(("http://", "https://", "data:")):
                return match.group(0)
            asset_path = (base_dir / original).resolve()
            if not asset_path.exists() or not asset_path.is_file():
                return match.group(0)
            mime_type = "image/png"
            if asset_path.suffix.lower() in {".jpg", ".jpeg"}:
                mime_type = "image/jpeg"
            elif asset_path.suffix.lower() == ".webp":
                mime_type = "image/webp"
            elif asset_path.suffix.lower() == ".svg":
                mime_type = "image/svg+xml"
            data_url = f"data:{mime_type};base64,{base64.b64encode(asset_path.read_bytes()).decode('ascii')}"
            return match.group(0).replace(original, data_url)

        return re.sub(r'src="([^"]+)"', replace_src, html)

    def _prepare_slide_pdf_html(self, html: str) -> str:
        if "<body>" in html:
            return html.replace("<body>", '<body class="pdf-export">', 1)
        return html.replace("<body ", '<body class="pdf-export" ', 1)

    def _short_support_copy(self, detail: str) -> str:
        words = detail.split()
        if len(words) <= 12:
            return detail
        return " ".join(words[:12]).rstrip(" ,.;:") + "."

    def _format_inline_html(self, text: str) -> str:
        escaped = self._escape_html(text)
        return re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", escaped)

    def _build_presentation_pdf(self, markdown_path: Path, html_path: Path, pdf_path: Path, slides: list[Slide]) -> str:
        html_text = self._prepare_slide_pdf_html(self._inline_slide_assets(html_path.read_text(encoding="utf-8"), html_path.parent))
        try:
            with sync_playwright() as playwright:
                browser = playwright.chromium.launch(headless=True, args=["--no-sandbox", "--disable-dev-shm-usage"])
                page = browser.new_page(viewport={"width": 1920, "height": 1080}, device_scale_factor=1)
                page.set_content(html_text, wait_until="networkidle")
                page.emulate_media(media="print")
                page.pdf(
                    path=str(pdf_path),
                    width="1920px",
                    height="1080px",
                    print_background=True,
                    margin={"top": "0", "right": "0", "bottom": "0", "left": "0"},
                )
                browser.close()
            return "playwright-html"
        except Exception:
            provider = self._render_slide_pdf(markdown_path, html_path, pdf_path, html_text)
            if pdf_path.exists():
                return provider
            self._build_slide_pdf_with_reportlab(pdf_path, slides)
            return "reportlab-slides"

    def _build_slide_pdf_with_reportlab(self, path: Path, slides: list[Slide]) -> None:
        width, height = 13.333 * inch, 7.5 * inch
        pdf = pdf_canvas.Canvas(str(path), pagesize=(width, height))
        for index, slide in enumerate(slides):
            self._draw_slide_page(pdf, slide, index, len(slides), width, height)
            pdf.showPage()
        pdf.save()

    def _draw_slide_page(self, pdf: Any, slide: Slide, index: int, total: int, width: float, height: float) -> None:
        dark = slide.layout in {"lead", "closing"}
        bg = colors.HexColor("#123657") if dark else colors.HexColor("#f3f8fd")
        accent = colors.HexColor("#27a6b8")
        ink = colors.white if dark else colors.HexColor("#0c2848")
        muted = colors.HexColor("#d7e8f7") if dark else colors.HexColor("#49637f")
        pdf.setFillColor(bg)
        pdf.rect(0, 0, width, height, fill=1, stroke=0)
        pdf.setFillColor(colors.HexColor("#1e5a84") if dark else colors.HexColor("#dfeaf4"))
        pdf.rect(width * 0.56, 0, width * 0.44, height, fill=1, stroke=0)
        pdf.setFillColor(accent)
        pdf.circle(width * 0.86, height * 0.25, 90, fill=1, stroke=0)
        pdf.setFillColor(colors.Color(1, 1, 1, alpha=0.12) if dark else colors.Color(1, 1, 1, alpha=0.62))
        pdf.circle(width * 0.84, height * 0.68, 145, fill=1, stroke=0)

        left = 0.62 * inch
        top = height - 0.68 * inch
        pdf.setFillColor(muted)
        pdf.setFont("Helvetica-Bold", 9)
        pdf.drawString(left, top, (slide.kicker or f"Slide {index + 1}").upper()[:72])
        pdf.drawRightString(width - 0.62 * inch, top, f"{index + 1} / {total}")

        title_size = 42 if len(slide.title) <= 34 else 34
        title_lines = self._wrap_text(slide.title, 24 if title_size >= 40 else 32, max_lines=3)
        y = top - 0.58 * inch
        pdf.setFillColor(ink)
        pdf.setFont("Helvetica-Bold", title_size)
        for line in title_lines:
            pdf.drawString(left, y, line)
            y -= title_size * 1.06

        if slide.layout == "lead" and slide.body:
            pdf.setFillColor(muted)
            pdf.setFont("Helvetica", 18)
            for line in self._wrap_text(slide.body[0], 34, max_lines=2):
                pdf.drawString(left, y - 8, line)
                y -= 24
            y -= 18

        items = (slide.bullets or slide.body or [])[:4]
        if slide.layout in {"highlights", "agenda", "metrics"} and items:
            self._draw_slide_cards(pdf, items, left, 0.82 * inch, width * 0.50, max(y - 0.2 * inch, 2.8 * inch), dark)
        elif items:
            pdf.setFont("Helvetica", 16)
            pdf.setFillColor(ink)
            for item in items:
                label, detail = self._split_card_item(item)
                text = detail if detail and detail != "Ponto principal do slide" else label
                for line_index, line in enumerate(self._wrap_text(text, 58, max_lines=2)):
                    if line_index == 0:
                        pdf.setFillColor(accent)
                        pdf.circle(left + 7, y - 5, 4, fill=1, stroke=0)
                        pdf.setFillColor(ink)
                    pdf.drawString(left + 24, y - 12, line)
                    y -= 22
                y -= 10

    def _draw_slide_cards(self, pdf: Any, items: list[str], left: float, bottom: float, width: float, top: float, dark: bool) -> None:
        card_gap = 0.22 * inch
        card_w = (width - card_gap) / 2
        card_h = max(1.08 * inch, (top - bottom - card_gap) / 2)
        ink = colors.HexColor("#0c2848")
        muted = colors.HexColor("#49637f")
        accent = colors.HexColor("#1f5ea8")
        for idx, item in enumerate(items[:4]):
            col = idx % 2
            row = idx // 2
            x = left + col * (card_w + card_gap)
            y = top - (row + 1) * card_h - row * card_gap
            pdf.setFillColor(colors.white if not dark else colors.HexColor("#f7fbff"))
            pdf.roundRect(x, y, card_w, card_h, 14, fill=1, stroke=0)
            label, detail = self._split_card_item(item)
            pdf.setFillColor(accent)
            pdf.setFont("Helvetica-Bold", 8)
            pdf.drawString(x + 18, y + card_h - 24, label.upper()[:34])
            pdf.setFillColor(ink)
            pdf.setFont("Helvetica-Bold", 18)
            text = detail if detail and detail != "Ponto principal do slide" else label
            text_lines = self._wrap_text(text, 22, max_lines=3)
            line_y = y + card_h - 52
            for line in text_lines:
                pdf.drawString(x + 18, line_y, line)
                line_y -= 21
            pdf.setFillColor(muted)
            pdf.setFont("Helvetica", 9)
            support = self._short_support_copy(text)
            for line in self._wrap_text(support, 30, max_lines=2):
                pdf.drawString(x + 18, max(y + 18, line_y - 4), line)
                line_y -= 12

    def _wrap_text(self, text: str, max_chars: int, max_lines: int = 4) -> list[str]:
        words = self._clean_inline_markdown(text).split()
        lines: list[str] = []
        current: list[str] = []
        for word in words:
            candidate = " ".join(current + [word])
            if current and len(candidate) > max_chars:
                lines.append(" ".join(current))
                current = [word]
                if len(lines) >= max_lines:
                    break
            else:
                current.append(word)
        if current and len(lines) < max_lines:
            lines.append(" ".join(current))
        if len(lines) == max_lines and len(" ".join(words)) > len(" ".join(lines)):
            lines[-1] = lines[-1].rstrip(".,;:") + "."
        return lines or [""]

    def _infer_slide_class(self, slide: str, index: int, total: int) -> str:
        lower = slide.lower()
        if index == 0:
            return "lead"
        if index == total - 1:
            return "closing"
        if "|" in slide and "\n" in slide:
            return "compare"
        if any(token in lower for token in ["agenda", "roteiro", "sumario", "sumário", "visao geral", "visão geral"]):
            return "agenda"
        if any(token in lower for token in ["roadmap", "cronograma", "etapas", "fases", "processo", "passos", "plano"]):
            return "timeline"
        bullet_count = len(re.findall(r"(?m)^[-*]\s+", slide))
        numeric_hits = len(re.findall(r"\b\d+(?:[%x]|(?:[.,]\d+)?)\b", slide))
        if bullet_count >= 3 and numeric_hits >= 2:
            return "metrics"
        card_like_bullets = [
            line
            for line in slide.splitlines()
            if re.match(r"^[-*]\s+.+(?:\s[-:]\s+|\s:\s+).+", line.strip())
        ]
        if len(card_like_bullets) >= 3:
            return "highlights"
        return "section"

    def _inject_slide_class(self, slide: str, slide_class: str) -> str:
        directive = f"<!-- _class: {slide_class} -->"
        if slide.lstrip().startswith("<!-- _class:"):
            return slide
        return f"{directive}\n\n{slide.strip()}"

    def _normalize_agenda_slide(self, slide: str) -> str:
        lines = [line.strip() for line in slide.splitlines() if line.strip()]
        if not lines:
            return slide
        heading = lines[0]
        body = []
        for line in lines[1:]:
            if re.match(r"^[-*]\s+", line):
                body.append(line)
            elif not line.startswith(">"):
                body.append(f"- {line}")
        if not body:
            return slide
        return "\n".join([heading, ""] + body)

    def _render_metrics_slide(self, slide: str) -> str:
        title, bullets, extras = self._split_slide_content(slide)
        cards: list[str] = []
        for bullet in bullets[:4]:
            label, value, text = self._parse_metric_bullet(bullet)
            cards.append(
                '<div class="metric-card">'
                f'<span class="metric-label">{self._escape_html(label)}</span>'
                f'<span class="metric-value">{self._escape_html(value)}</span>'
                f'<span class="metric-text">{self._escape_html(text)}</span>'
                "</div>"
            )
        html_block = '<div class="metrics-grid">' + "".join(cards) + "</div>"
        parts = [title, "", html_block]
        if extras:
            parts.extend([""] + extras)
        return "\n".join(parts)

    def _render_timeline_slide(self, slide: str) -> str:
        title, bullets, extras = self._split_slide_content(slide)
        if not bullets:
            return slide
        ordered = [re.sub(r"^[-*]\s+", "", bullet).strip() for bullet in bullets[:5]]
        body = [title, "", "<ol>"] + [f"<li>{self._escape_html(item)}</li>" for item in ordered] + ["</ol>"]
        if extras:
            body.extend([""] + extras)
        return "\n".join(body)

    def _render_closing_slide(self, slide: str, title: str) -> str:
        lines = [line.strip() for line in slide.splitlines() if line.strip()]
        if not lines:
            return slide
        if len(lines) == 1:
            return f"# {title}\n\n## Recomendacao final\n\n- Priorize execucao simples, narrativa clara e proximos passos concretos."
        if not any(line.startswith("## ") for line in lines[1:]):
            lines.insert(1, "## Recomendacao final")
        return "\n".join(lines)

    def _split_slide_content(self, slide: str) -> tuple[str, list[str], list[str]]:
        lines = [line.strip() for line in slide.splitlines() if line.strip() and not line.strip().startswith("<!--")]
        if not lines:
            return "## Slide", [], []
        title = lines[0]
        bullets = [line for line in lines[1:] if re.match(r"^[-*]\s+", line)]
        extras = [line for line in lines[1:] if line not in bullets]
        return title, bullets, extras

    def _parse_metric_bullet(self, bullet: str) -> tuple[str, str, str]:
        content = re.sub(r"^[-*]\s+", "", bullet).strip()
        strong_match = re.match(r"^\*\*(.+?)\*\*\s*[-:]\s*(.+)$", content)
        if strong_match:
            label = strong_match.group(1).strip()
            tail = strong_match.group(2).strip()
        else:
            parts = re.split(r"\s[-:]\s", content, maxsplit=1)
            label = parts[0].strip()
            tail = parts[1].strip() if len(parts) > 1 else ""
        value_match = re.match(r"^([0-9][^,;.]{0,24})(?:\s*[,-]\s*|\s+)(.+)$", tail)
        if value_match:
            value = value_match.group(1).strip()
            text = value_match.group(2).strip()
        else:
            tokens = tail.split()
            value = tokens[0] if tokens else label
            text = " ".join(tokens[1:]).strip() or tail or "Indicador-chave da narrativa."
        return label or "Indicador", value, text

    def _slide_title_from_text(self, text: str) -> str:
        for line in text.splitlines():
            cleaned = re.sub(r"^#{1,3}\s*", "", line).strip()
            if cleaned:
                return cleaned[:80]
        return "Apresentacao"

    def _normalize_slide(self, section: str) -> str:
        lines = [line.rstrip() for line in section.splitlines() if line.strip()]
        if not lines:
            return "## Slide\n\n- Conteudo"
        if not re.match(r"^#{1,2}\s+", lines[0]):
            lines.insert(0, "## Slide")
        return "\n".join(lines)

    def _upgrade_cover_slide(self, slide: str) -> str:
        lines = [line for line in slide.splitlines() if line.strip()]
        if not lines:
            return slide
        title = re.sub(r"^#{1,2}\s+", "", lines[0]).strip()
        body = [line for line in lines[1:] if line.strip()]
        subtitle = body[0] if body else "Visao geral executiva"
        return (
            f"# {title}\n\n"
            f"## {subtitle}\n\n"
            "> Panorama claro, visual e pronto para apresentar.\n"
        )

    def _slide_bullets(self, text: str) -> str:
        chunks = [chunk.strip() for chunk in re.split(r"(?<=[\.\!\?])\s+", " ".join(text.split())) if chunk.strip()]
        bullets = [f"- {chunk}" for chunk in chunks[:5]]
        return "\n".join(bullets) or "- Conteudo principal"

    def _render_marp_html(self, markdown_path: Path, html_path: Path) -> bool:
        command = self._marp_command(markdown_path, html_path, "--html")
        try:
            subprocess.run(command, check=True, capture_output=True, text=True)
            return True
        except Exception:
            return False

    def _render_slide_pdf(self, markdown_path: Path, html_path: Path, pdf_path: Path, html_text: str | None = None) -> str:
        html_payload = html_text or self._inline_slide_assets(html_path.read_text(encoding="utf-8"), html_path.parent)
        if self.settings.gotenberg_url and html_path.exists():
            try:
                self._build_pdf_with_gotenberg(pdf_path, markdown_path.stem, html_payload, slide_layout=True)
                return "presentation-html + gotenberg"
            except Exception:
                pass
        return "presentation-html-unavailable"

    def _marp_command(self, input_path: Path, output_path: Path, mode: str) -> list[str]:
        marp_bin = shutil.which("marp")
        if marp_bin:
            return [marp_bin, str(input_path), mode, "-o", str(output_path), "--allow-local-files"]
        npx_bin = shutil.which("npx") or shutil.which("npx.cmd")
        if npx_bin:
            return [npx_bin, "--yes", "@marp-team/marp-cli", str(input_path), mode, "-o", str(output_path), "--allow-local-files"]
        raise FileNotFoundError("Marp CLI nao encontrado.")

    def _filename_stem(self, user_request: str, title: str) -> str:
        request_text = " ".join(user_request.split()).strip(" .:-")
        if request_text:
            return request_text[:120]
        return title

    def _safe_filename(self, title: str, extension: str) -> str:
        slug = re.sub(r"[^a-zA-Z0-9]+", "-", title.lower()).strip("-")
        return f"{slug or 'documento-kimi-ai'}{extension}"

    def _normalize_document_text(self, text: str, user_request: str) -> str:
        cleaned = text.strip()
        if not cleaned:
            return cleaned
        cleaned = cleaned.replace("\u25a0", " ")
        cleaned = cleaned.replace("\u00a0", " ")
        cleaned = re.sub(r"[\u200b-\u200f\u202a-\u202e]", "", cleaned)
        cleaned = re.sub(r"^```(?:markdown|md|text|python|py|html|javascript|js|typescript|ts)?\s*", "", cleaned, flags=re.I)
        cleaned = re.sub(r"\s*```$", "", cleaned)
        cleaned = re.sub(
            r"(?is)^#+\s*pedido:.*?(?=^#+\s+|^[A-ZÁÀÂÃÉÈÊÍÌÎÓÒÔÕÚÙÛÇa-záàâãéèêíìîóòôõúùûç].*$|$)",
            "",
            cleaned,
        ).strip()
        lead_patterns = [
            r"(?is)^compreendido\.?\s+voc[eê].{0,220}?(?=\n\s*\n|^#|\Z)",
            r"(?is)^aqui est[aá].{0,220}?(?=\n\s*\n|^#|\Z)",
            r"(?is)^abaixo est[aá].{0,220}?(?=\n\s*\n|^#|\Z)",
            r"(?is)^segue.{0,180}?(?=\n\s*\n|^#|\Z)",
            r"(?is)^este documento.{0,220}?(?=\n\s*\n|^#|\Z)",
            r"(?is)^o documento a seguir.{0,220}?(?=\n\s*\n|^#|\Z)",
        ]
        for pattern in lead_patterns:
            cleaned = re.sub(pattern, "", cleaned).strip()
        request_topic = self._extract_topic(user_request).lower()
        cleaned = re.sub(
            r"(?is)^pedido original:\s*.*?(?=\n\s*\n|^#|\Z)",
            "",
            cleaned,
        ).strip()
        if request_topic:
            cleaned = re.sub(
                r"(?is)^int[eê]n[cç][aã]o:.*?(?=\n\s*\n|^#|\Z)",
                "",
                cleaned,
            ).strip()
        return cleaned

    def _strip_system_notices(self, text: str) -> str:
        lines = [line.rstrip() for line in text.splitlines()]
        cleaned: list[str] = []
        skipping_notice = False
        for line in lines:
            stripped = line.strip()
            if stripped.lower().startswith("aviso de sistema:"):
                skipping_notice = True
                continue
            if skipping_notice and not stripped:
                skipping_notice = False
                continue
            if skipping_notice and re.search(r"\b(gemini|groq|cerebras|openrouter|fallback|indisponivel|indisponível)\b", stripped, re.I):
                continue
            cleaned.append(line)
        result = "\n".join(cleaned).strip()
        return result or text.strip()

    def _looks_like_document_content(self, text: str) -> bool:
        if not text:
            return False
        compact = " ".join(text.split())
        if len(compact) < 120:
            return False
        lowered = compact.lower()
        banned_markers = [
            "aviso de sistema:",
            "fallback ativo",
            "indisponivel para esta chave",
            "indisponível para esta chave",
            "classificacao interna",
            "intenção:",
            "intencao:",
        ]
        if any(marker in lowered for marker in banned_markers):
            return False
        if self._looks_like_code_dump(text):
            return False
        return True

    def _looks_like_code_dump(self, text: str) -> bool:
        lowered = text.lower()
        code_markers = [
            "from docx import",
            "import docx",
            "def ",
            "class ",
            "paragraph.",
            "run.font",
            "wd_align_paragraph",
            "oxmlelement",
            "qn(",
            "cm(",
            "pt(",
            "python-docx",
            "document()",
            "add_heading(",
            "add_paragraph(",
            "```python",
            "```py",
        ]
        marker_hits = sum(1 for marker in code_markers if marker in lowered)
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        code_like_lines = sum(
            1
            for line in lines
            if re.match(r"^(def|class|from|import|for|if|elif|else|return|with)\b", line)
            or re.search(r"[A-Za-z_][A-Za-z0-9_]*\s*=", line)
            or line.endswith(":")
        )
        if marker_hits >= 2:
            return True
        return bool(lines) and code_like_lines / max(len(lines), 1) > 0.28

    def _fallback_document_text(self, user_request: str) -> str:
        topic = self._extract_topic(user_request)
        return (
            f"# {topic}\n\n"
            "## Resumo executivo\n"
            f"Este documento apresenta uma visao organizada sobre {topic}, com foco em contexto, principais acontecimentos, impactos e pontos de atencao.\n\n"
            "## Contexto geral\n"
            f"{topic} deve ser entendido a partir do seu contexto historico, politico, social e tecnologico. "
            "O objetivo aqui e reunir uma base clara para leitura, estudo ou apresentacao.\n\n"
            "## Principais pontos\n"
            f"- Origem e antecedentes relacionados a {topic}\n"
            f"- Eventos, marcos ou fases centrais de {topic}\n"
            f"- Consequencias diretas e indiretas de {topic}\n"
            f"- Leitura critica sobre os impactos de {topic}\n\n"
            "## Desenvolvimento\n"
            f"A analise de {topic} pode ser organizada em uma linha do tempo, nos atores envolvidos e nos efeitos produzidos ao longo do tempo. "
            "Em um documento final, essa secao serve para aprofundar os fatos e organizar a narrativa de forma legivel.\n\n"
            "## Impactos e interpretacoes\n"
            f"Os impactos de {topic} podem incluir transformacoes politicas, economicas, culturais, sociais ou estrategicas. "
            "Tambem e importante destacar controversias, diferentes interpretacoes e consequencias de longo prazo.\n\n"
            "## Conclusao\n"
            f"{topic} e um tema que exige organizacao clara, criterio e boa hierarquia textual. "
            "Este PDF e DOCX foram estruturados para servir como ponto de partida para leitura, estudo e refinamento posterior.\n"
        )

    def _extract_topic(self, user_request: str) -> str:
        text = " ".join(user_request.split()).strip()
        text = re.sub(r"^(me\s+)?(gere|gerar|gera|crie|criar|fa[cç]a|fazer|monte|montar|produza|produzir)\s+", "", text, flags=re.I)
        text = re.sub(r"^(um|uma|o|a)\s+", "", text, flags=re.I)
        text = re.sub(r"\b(pdf|docx|docxs|docs?|arquivo|documento|word|markdown|slide|slides|deck|pptx?|apresenta[cç][aã]o)\b", "", text, flags=re.I)
        sobre_match = re.search(r"\bsobre\s+(.+)$", text, re.I)
        if sobre_match and sobre_match.group(1).strip(" .:-"):
            topic = re.split(
                r"\b(citando|incluindo|inclua|usando|use|no estilo|em estilo|formato|de forma|nivel|nível)\b",
                sobre_match.group(1),
                maxsplit=1,
                flags=re.I,
            )[0].strip(" .:-")
            if topic:
                return self._polish_topic_title(topic)
        intent_match = re.search(
            r"\bpara\s+(?:anunciar|divulgar|vender|promover|explicar|ensinar|apresentar)\s+(?:um|uma|o|a)?\s*(.+)$",
            text,
            re.I,
        )
        if intent_match and intent_match.group(1).strip(" .:-"):
            return self._polish_topic_title(intent_match.group(1).strip(" .:-"))
        text = re.split(
            r"\b(citando|com|incluindo|inclua|usando|use|no estilo|em estilo|formato|de forma|para|nivel|nível)\b",
            text,
            maxsplit=1,
            flags=re.I,
        )[0]
        text = re.sub(r"\b(suas|devidas|devida|devido)\b", "", text, flags=re.I)
        text = re.sub(r"\bisos?\b", "ISO", text, flags=re.I)
        text = re.sub(r"\s{2,}", " ", text).strip(" .:-")
        patterns = [
            r"(?:sobre|do|da|de)\s+(.+)$",
            r"(?:pedido|tema)\s+(.+)$",
        ]
        for pattern in patterns:
            match = re.search(pattern, text, re.I)
            if match:
                topic = match.group(1).strip(" .:-")
                if topic:
                    return self._polish_topic_title(topic)
        return self._polish_topic_title(text[:90] or "Documento solicitado")

    def _polish_topic_title(self, topic: str) -> str:
        cleaned = self._clean_inline_markdown(topic).strip(" .:-")
        replacements = {
            "gestao de ativos na tecnologia da informacao": "Gestao de Ativos de TI",
            "gestão de ativos na tecnologia da informação": "Gestao de Ativos de TI",
            "gestao de ativos em ti": "Gestao de Ativos de TI",
            "gestão de ativos em ti": "Gestao de Ativos de TI",
        }
        lowered = cleaned.lower()
        for needle, replacement in replacements.items():
            if needle in lowered:
                return replacement
        small_words = {"de", "da", "do", "das", "dos", "e", "em", "na", "no", "para", "com"}
        words: list[str] = []
        for word in cleaned.split():
            if word.upper() in {"TI", "ISO", "ISOS", "LGPD", "IA"}:
                words.append(word.upper().replace("ISOS", "ISO"))
            elif word.lower() in small_words:
                words.append(word.lower())
            else:
                words.append(word[:1].upper() + word[1:].lower())
        return " ".join(words)[:80] or "Apresentacao"

    def _parse_blocks(self, text: str) -> list[Block]:
        blocks: list[Block] = []
        bullet_buffer: list[str] = []
        numbered_buffer: list[str] = []
        table_buffer: list[str] = []

        def flush_buffers() -> None:
            nonlocal bullet_buffer, numbered_buffer, table_buffer
            if table_buffer:
                rows = self._parse_markdown_table(table_buffer)
                if rows:
                    blocks.append(Block("table", "", rows=rows))
                else:
                    for item in table_buffer:
                        if item.strip():
                            blocks.append(Block("paragraph", self._clean_inline_markdown(item.strip())))
            for item in bullet_buffer:
                blocks.append(Block("bullet", self._clean_inline_markdown(item)))
            for item in numbered_buffer:
                blocks.append(Block("numbered", self._clean_inline_markdown(item)))
            table_buffer = []
            bullet_buffer = []
            numbered_buffer = []

        for raw_line in text.splitlines():
            line = raw_line.strip()
            if not line:
                flush_buffers()
                continue
            if line.startswith("```"):
                continue
            if re.fullmatch(r"-{3,}", line):
                continue
            if self._looks_like_table_line(line):
                table_buffer.append(line)
                continue
            if table_buffer:
                flush_buffers()
            heading = re.match(r"^(#{1,3})\s+(.*)$", line)
            if heading:
                flush_buffers()
                blocks.append(Block("heading", self._clean_inline_markdown(heading.group(2).strip()), len(heading.group(1))))
                continue
            bullet = re.match(r"^[-*]\s+(.*)$", line)
            if bullet:
                bullet_buffer.append(bullet.group(1).strip())
                continue
            numbered = re.match(r"^\d+[.)]\s+(.*)$", line)
            if numbered:
                numbered_buffer.append(numbered.group(1).strip())
                continue
            flush_buffers()
            blocks.append(Block("paragraph", self._clean_inline_markdown(line)))

        flush_buffers()
        if not blocks:
            blocks.append(Block("paragraph", text.strip() or "Documento sem conteudo."))
        return blocks

    def _build_docx(self, path: Path, title: str, blocks: list[Block], user_request: str) -> None:
        doc = Document()
        section = doc.sections[0]
        section.top_margin = DocxInches(1)
        section.bottom_margin = DocxInches(1)
        section.left_margin = DocxInches(1)
        section.right_margin = DocxInches(1)

        self._configure_styles(doc)
        self._configure_header(section, title)
        self._configure_footer(section)

        doc.add_paragraph(title, style="Title")
        doc.add_paragraph()

        for block in blocks:
            if block.kind == "heading":
                style_name = {1: "Heading 1", 2: "Heading 2", 3: "Heading 3"}.get(block.level, "Heading 2")
                doc.add_paragraph(block.text, style=style_name)
            elif block.kind == "table" and block.rows:
                self._append_docx_table(doc, block.rows)
            elif block.kind == "bullet":
                doc.add_paragraph(block.text, style="List Bullet")
            elif block.kind == "numbered":
                doc.add_paragraph(block.text, style="List Number")
            else:
                doc.add_paragraph(block.text, style="Normal")

        doc.save(path)

    def _build_pdf(self, path: Path, title: str, blocks: list[Block], user_request: str, session_id: str) -> str:
        if self._is_slide_request(user_request):
            source_text = "\n".join(
                f"## {block.text}" if block.kind == "heading" else f"- {block.text}" if block.kind in {"bullet", "numbered"} else block.text
                for block in blocks
                if block.text
            )
            normalized_source = self._fallback_slide_deck_text(user_request, title) if self._needs_slide_fallback(source_text, user_request) else self._normalize_slide_source(source_text)
            raw_slides = self._split_raw_slides(normalized_source) if self._looks_like_marp_deck(normalized_source) else self._split_into_slides(normalized_source)
            if len(raw_slides) < 6:
                raw_slides = self._split_into_slides(self._fallback_slide_deck_text(user_request, title))
            slides = self._build_slide_models(raw_slides, title, {})
            html_path = path.with_suffix(".html")
            markdown_path = path.with_suffix(".md")
            html_path.write_text(self._build_presentation_html(title, slides, user_request), encoding="utf-8")
            markdown_path.write_text(self._build_marp_markdown(title, normalized_source), encoding="utf-8")
            return self._build_presentation_pdf(markdown_path, html_path, path, slides)
        html = self._build_html(title, blocks, user_request, session_id)
        if self.settings.gotenberg_url:
            try:
                self._build_pdf_with_gotenberg(path, title, html)
                return "gotenberg"
            except Exception:
                pass
        try:
            self._build_pdf_with_playwright(path, html)
            return "playwright"
        except Exception:
            pass

        self._build_pdf_with_reportlab(path, title, blocks, user_request)
        return "reportlab"

    def _build_pdf_with_reportlab(self, path: Path, title: str, blocks: list[Block], user_request: str) -> None:
        if self._is_slide_request(user_request):
            slide_blocks = self._group_slides(title, blocks)
            slides = [
                Slide(
                    title=item["title"],
                    body=[entry["text"] for entry in item["items"] if entry["kind"] != "bullet"] or None,
                    bullets=[entry["text"] for entry in item["items"] if entry["kind"] == "bullet"] or None,
                    kicker=f"Slide {index + 1}",
                )
                for index, item in enumerate(slide_blocks)
            ]
            self._build_slide_pdf_with_reportlab(path, slides)
            return
        styles = getSampleStyleSheet()
        body = ParagraphStyle(
            "KimiBody",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=11,
            leading=16,
            spaceAfter=10,
            textColor=colors.HexColor("#243240"),
        )
        body.wordWrap = "CJK"
        h1 = ParagraphStyle("KimiH1", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=16, leading=20, textColor=colors.HexColor("#14213d"), spaceBefore=10, spaceAfter=8)
        h2 = ParagraphStyle("KimiH2", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=13, leading=17, textColor=colors.HexColor("#223b63"), spaceBefore=8, spaceAfter=6)
        title_style = ParagraphStyle("KimiTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=22, leading=26, textColor=colors.HexColor("#0f172a"), spaceAfter=8)
        table_cell = ParagraphStyle("KimiTableCell", parent=body, fontSize=10, leading=14, spaceAfter=0)
        table_header = ParagraphStyle("KimiTableHeader", parent=table_cell, fontName="Helvetica-Bold", textColor=colors.white)
        story: list[Any] = [Paragraph(title, title_style), Spacer(1, 0.08 * inch)]

        bullet_items: list[ListItem] = []
        numbered_items: list[ListItem] = []

        def flush_lists() -> None:
            nonlocal bullet_items, numbered_items
            if bullet_items:
                story.append(ListFlowable(bullet_items, bulletType="bullet", leftIndent=18))
                story.append(Spacer(1, 0.08 * inch))
                bullet_items = []
            if numbered_items:
                story.append(ListFlowable(numbered_items, bulletType="1", leftIndent=18))
                story.append(Spacer(1, 0.08 * inch))
                numbered_items = []

        for block in blocks:
            if block.kind == "bullet":
                bullet_items.append(ListItem(Paragraph(block.text, body)))
                continue
            if block.kind == "numbered":
                numbered_items.append(ListItem(Paragraph(block.text, body)))
                continue
            flush_lists()
            if block.kind == "heading":
                story.append(Paragraph(block.text, h1 if block.level == 1 else h2))
            elif block.kind == "table" and block.rows:
                story.append(self._build_pdf_table(block.rows, table_header, table_cell))
                story.append(Spacer(1, 0.12 * inch))
            else:
                story.append(Paragraph(block.text, body))
        flush_lists()

        pdf = SimpleDocTemplate(str(path), pagesize=LETTER, leftMargin=inch, rightMargin=inch, topMargin=inch, bottomMargin=inch)
        pdf.build(story)

    def _build_slide_pdf_with_reportlab_legacy(self, path: Path, title: str, blocks: list[Block]) -> None:
        styles = getSampleStyleSheet()
        slide_title = ParagraphStyle(
            "KimiSlideTitle",
            parent=styles["Title"],
            fontName="Helvetica-Bold",
            fontSize=24,
            leading=28,
            textColor=colors.HexColor("#0f172a"),
            spaceAfter=14,
        )
        slide_body = ParagraphStyle(
            "KimiSlideBody",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=14,
            leading=20,
            textColor=colors.HexColor("#334155"),
            spaceAfter=10,
        )
        slides = self._group_slides(title, blocks)
        story: list[Any] = []
        for index, slide in enumerate(slides):
            story.append(Paragraph(slide["title"], slide_title))
            for item in slide["items"]:
                prefix = "• " if item["kind"] == "bullet" else ""
                story.append(Paragraph(f"{prefix}{item['text']}", slide_body))
            if index < len(slides) - 1:
                from reportlab.platypus import PageBreak

                story.append(PageBreak())

        pdf = SimpleDocTemplate(
            str(path),
            pagesize=landscape(LETTER),
            leftMargin=0.75 * inch,
            rightMargin=0.75 * inch,
            topMargin=0.65 * inch,
            bottomMargin=0.65 * inch,
        )
        pdf.build(story)

    def _build_pdf_with_gotenberg(self, path: Path, title: str, html: str, slide_layout: bool = False) -> None:
        base_url = (self.settings.gotenberg_url or "").rstrip("/")
        if not base_url:
            raise ValueError("Gotenberg URL nao configurada.")
        if "://" not in base_url:
            base_url = f"http://{base_url}"
        headers = {"Gotenberg-Output-Filename": self._safe_filename(title, "").strip(".-") or "documento-kimi-ai"}
        data = {
            "printBackground": "true",
            "generateDocumentOutline": "true",
        }
        if slide_layout:
            data.update(
                {
                    "generateDocumentOutline": "false",
                    "preferCssPageSize": "true",
                    "paperWidth": "13.333",
                    "paperHeight": "7.5",
                    "marginTop": "0",
                    "marginBottom": "0",
                    "marginLeft": "0",
                    "marginRight": "0",
                    "emulatedMediaType": "screen",
                    "scale": "1.0",
                }
            )
        files = [("files", ("index.html", html.encode("utf-8"), "text/html; charset=utf-8"))]
        with httpx.Client(timeout=self.settings.gotenberg_timeout_seconds) as client:
            response = client.post(f"{base_url}/forms/chromium/convert/html", headers=headers, data=data, files=files)
            response.raise_for_status()
        path.write_bytes(response.content)

    def _build_pdf_with_playwright(self, path: Path, html: str) -> None:
        try:
            with sync_playwright() as playwright:
                browser = playwright.chromium.launch(headless=True, args=["--no-sandbox", "--disable-dev-shm-usage"])
                page = browser.new_page()
                page.set_content(html, wait_until="networkidle")
                page.pdf(
                    path=str(path),
                    format="A4",
                    print_background=True,
                    margin={
                        "top": "24mm",
                        "right": "18mm",
                        "bottom": "22mm",
                        "left": "18mm",
                    },
                )
                browser.close()
        except PlaywrightError as exc:
            raise RuntimeError("Falha ao gerar PDF com Playwright.") from exc

    def _build_html(self, title: str, blocks: list[Block], user_request: str, session_id: str) -> str:
        if self._is_slide_request(user_request):
            return self._build_slide_html(title, blocks)
        content_parts = []
        open_list: str | None = None
        for block in blocks:
            if block.kind in {"bullet", "numbered"}:
                tag = "ul" if block.kind == "bullet" else "ol"
                if open_list != tag:
                    if open_list:
                        content_parts.append(f"</{open_list}>")
                    content_parts.append(f"<{tag}>")
                    open_list = tag
                content_parts.append(f"<li>{self._escape_html(block.text)}</li>")
                continue
            if open_list:
                content_parts.append(f"</{open_list}>")
                open_list = None
            if block.kind == "heading":
                level = min(max(block.level, 1), 3)
                content_parts.append(f"<h{level}>{self._escape_html(block.text)}</h{level}>")
            elif block.kind == "table" and block.rows:
                content_parts.append(self._build_html_table(block.rows))
            else:
                content_parts.append(f"<p>{self._escape_html(block.text)}</p>")
        if open_list:
            content_parts.append(f"</{open_list}>")

        return f"""<!doctype html>
<html lang="pt-BR">
<head>
  <meta charset="utf-8" />
  <title>{self._escape_html(title)}</title>
  <style>
    @page {{
      size: A4;
      margin: 24mm 18mm 22mm 18mm;
    }}
    body {{
      font-family: Arial, Helvetica, sans-serif;
      color: #243240;
      margin: 0;
      font-size: 12px;
      line-height: 1.6;
      background: #ffffff;
    }}
    header {{
      margin-bottom: 20px;
    }}
    h1 {{
      font-size: 28px;
      line-height: 1.15;
      margin: 8px 0 6px;
      color: #0f172a;
    }}
    h2 {{
      font-size: 18px;
      margin: 24px 0 8px;
      color: #14213d;
    }}
    h3 {{
      font-size: 15px;
      margin: 18px 0 6px;
      color: #223b63;
    }}
    table {{
      width: 100%;
      border-collapse: collapse;
      margin: 0 0 14px;
      table-layout: fixed;
    }}
    th, td {{
      border: 1px solid #c8d4e3;
      padding: 8px 10px;
      text-align: left;
      vertical-align: top;
    }}
    th {{
      background: #1f4b8f;
      color: #ffffff;
      font-weight: 700;
    }}
    tbody tr:nth-child(even) {{
      background: #f4f7fb;
    }}
    p {{
      margin: 0 0 10px;
    }}
    ul, ol {{
      margin: 0 0 12px 22px;
      padding: 0;
    }}
    li {{
      margin: 0 0 8px;
    }}
  </style>
</head>
<body>
  <header>
    <h1>{self._escape_html(title)}</h1>
  </header>
  {''.join(content_parts)}
</body>
</html>"""

    def _build_slide_html(self, title: str, blocks: list[Block]) -> str:
        slides = self._group_slides(title, blocks)
        slides_markup = []
        for index, slide in enumerate(slides, start=1):
            items_markup = []
            for item in slide["items"]:
                tag = "li" if item["kind"] == "bullet" else "p"
                items_markup.append(f"<{tag}>{self._escape_html(item['text'])}</{tag}>")
            body_markup = "".join(items_markup) or "<p>Conteudo em preparacao.</p>"
            list_open = "<ul>" if any(item["kind"] == "bullet" for item in slide["items"]) else ""
            list_close = "</ul>" if any(item["kind"] == "bullet" for item in slide["items"]) else ""
            if list_open:
                bullets = "".join(f"<li>{self._escape_html(item['text'])}</li>" for item in slide["items"] if item["kind"] == "bullet")
                paragraphs = "".join(f"<p>{self._escape_html(item['text'])}</p>" for item in slide["items"] if item["kind"] != "bullet")
                body_markup = f"{paragraphs}{list_open}{bullets}{list_close}"
            slides_markup.append(
                f"""
  <section class="slide">
    <div class="slide-chrome">
      <span class="slide-index">{index:02d}</span>
      <span class="slide-mark">Kimi deck</span>
    </div>
    <div class="slide-body">
      <h1>{self._escape_html(slide["title"])}</h1>
      <div class="slide-content">{body_markup}</div>
    </div>
  </section>"""
            )
        return f"""<!doctype html>
<html lang="pt-BR">
<head>
  <meta charset="utf-8" />
  <title>{self._escape_html(title)}</title>
  <style>
    @page {{
      size: 13.333in 7.5in;
      margin: 0;
    }}
    * {{
      box-sizing: border-box;
    }}
    body {{
      margin: 0;
      font-family: Arial, Helvetica, sans-serif;
      background: #f4efe6;
      color: #111827;
    }}
    .slide {{
      width: 13.333in;
      height: 7.5in;
      padding: 0.5in;
      page-break-after: always;
      background:
        radial-gradient(circle at top left, rgba(52, 211, 153, 0.12), transparent 24%),
        linear-gradient(135deg, #fffdf8 0%, #f5efe5 100%);
    }}
    .slide:last-child {{
      page-break-after: auto;
    }}
    .slide-chrome {{
      display: flex;
      justify-content: space-between;
      align-items: center;
      margin-bottom: 0.35in;
      color: #6b7280;
      font-size: 11px;
      letter-spacing: 0.08em;
      text-transform: uppercase;
      font-weight: 700;
    }}
    .slide-body {{
      display: grid;
      gap: 0.26in;
      height: calc(100% - 0.7in);
      align-content: start;
      padding: 0.18in 0.08in;
    }}
    h1 {{
      margin: 0;
      font-size: 32px;
      line-height: 1.05;
      color: #0f172a;
    }}
    .slide-content {{
      display: grid;
      gap: 0.12in;
      font-size: 18px;
      line-height: 1.45;
      color: #334155;
    }}
    p {{
      margin: 0;
    }}
    ul {{
      margin: 0;
      padding-left: 0.24in;
    }}
    li {{
      margin: 0 0 0.08in;
    }}
  </style>
</head>
<body>
{''.join(slides_markup)}
</body>
</html>"""

    def _group_slides(self, title: str, blocks: list[Block]) -> list[dict[str, Any]]:
        slides: list[dict[str, Any]] = []
        current = {"title": title, "items": []}
        for block in blocks:
            if block.kind == "heading" and block.level <= 2:
                if current["items"]:
                    slides.append(current)
                current = {"title": block.text, "items": []}
                continue
            current["items"].append({"kind": "bullet" if block.kind == "bullet" else "text", "text": block.text})
            if len(current["items"]) >= 5:
                slides.append(current)
                current = {"title": title, "items": []}
        if current["items"] or not slides:
            slides.append(current)
        return slides

    def _is_slide_request_legacy(self, user_request: str) -> bool:
        lowered = " ".join(user_request.split()).lower()
        slide_markers = [
            "slide",
            "slides",
            "apresentacao",
            "apresentação",
            "deck",
            "pitch deck",
            "powerpoint",
            "ppt",
            "pptx",
        ]
        return any(marker in lowered for marker in slide_markers)

    def _escape_html(self, value: str) -> str:
        return (
            str(value)
            .replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
            .replace('"', "&quot;")
        )

    def _clean_inline_markdown(self, text: str) -> str:
        cleaned = str(text or "").replace("\u25a0", " ")
        cleaned = re.sub(r"[\u200b-\u200f\u202a-\u202e]", "", cleaned)
        cleaned = cleaned.replace("\\|", "|")
        cleaned = re.sub(r"\*\*(.*?)\*\*", r"\1", cleaned)
        cleaned = re.sub(r"__(.*?)__", r"\1", cleaned)
        cleaned = re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)", r"\1", cleaned)
        cleaned = re.sub(r"`([^`]+)`", r"\1", cleaned)
        cleaned = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1", cleaned)
        cleaned = re.sub(r"^#{1,6}\s*", "", cleaned)
        cleaned = re.sub(r"^>\s*", "", cleaned)
        cleaned = re.sub(r"\s+\|\s*$", "", cleaned)
        cleaned = re.sub(r"^\|\s*", "", cleaned)
        return " ".join(cleaned.split()).strip()

    def _looks_like_table_line(self, line: str) -> bool:
        stripped = line.strip()
        if stripped.count("|") < 2:
            return False
        if stripped.startswith("|") or stripped.endswith("|"):
            return True
        return bool(re.match(r"^[^|]+\|[^|]+\|", stripped))

    def _parse_markdown_table(self, lines: list[str]) -> list[list[str]]:
        rows: list[list[str]] = []
        for line in lines:
            stripped = line.strip()
            if not stripped:
                continue
            cells = [self._clean_inline_markdown(cell) for cell in stripped.strip("|").split("|")]
            if all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells):
                continue
            if any(cells):
                rows.append(cells)
        if len(rows) < 2:
            return []
        width = max(len(row) for row in rows)
        return [row + [""] * (width - len(row)) for row in rows]

    def _append_docx_table(self, doc: Document, rows: list[list[str]]) -> None:
        table = doc.add_table(rows=len(rows), cols=len(rows[0]))
        table.style = "Table Grid"
        table.autofit = True
        for row_index, row in enumerate(rows):
            for col_index, value in enumerate(row):
                cell = table.cell(row_index, col_index)
                cell.text = value
                if row_index == 0 and cell.paragraphs:
                    for run in cell.paragraphs[0].runs:
                        run.font.bold = True
        doc.add_paragraph()

    def _build_pdf_table(self, rows: list[list[str]], header_style: ParagraphStyle, cell_style: ParagraphStyle) -> Table:
        rendered = []
        for row_index, row in enumerate(rows):
            rendered.append(
                [
                    Paragraph(self._escape_reportlab(value), header_style if row_index == 0 else cell_style)
                    for value in row
                ]
            )
        table = Table(rendered, repeatRows=1, hAlign="LEFT")
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f4b8f")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("BACKGROUND", (0, 1), (-1, -1), colors.HexColor("#f8fafc")),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#ffffff"), colors.HexColor("#f4f7fb")]),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#c8d4e3")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 8),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                    ("TOPPADDING", (0, 0), (-1, -1), 6),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ]
            )
        )
        return table

    def _build_html_table(self, rows: list[list[str]]) -> str:
        header = "".join(f"<th>{self._escape_html(cell)}</th>" for cell in rows[0])
        body_rows = []
        for row in rows[1:]:
            body_rows.append("<tr>" + "".join(f"<td>{self._escape_html(cell)}</td>" for cell in row) + "</tr>")
        return (
            "<table><thead><tr>"
            + header
            + "</tr></thead><tbody>"
            + "".join(body_rows)
            + "</tbody></table>"
        )

    def _escape_reportlab(self, value: str) -> str:
        return self._escape_html(value).replace("\n", "<br/>")

    def _configure_styles(self, doc: Document) -> None:
        normal = doc.styles["Normal"]
        normal.font.name = "Arial"
        normal.font.size = Pt(11)

        title = doc.styles["Title"]
        title.font.name = "Arial"
        title.font.size = Pt(22)
        title.font.bold = True
        title.font.color.rgb = RGBColor(15, 23, 42)

        if "Subtitle" in doc.styles:
            subtitle = doc.styles["Subtitle"]
            subtitle.font.name = "Arial"
            subtitle.font.size = Pt(11)
            subtitle.font.color.rgb = RGBColor(91, 100, 114)

        for style_name, size in (("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)):
            style = doc.styles[style_name]
            style.font.name = "Arial"
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.color.rgb = RGBColor(20, 33, 61)

    def _configure_header(self, section: Any, label: str) -> None:
        paragraph = section.header.paragraphs[0]
        run = paragraph.add_run(label)
        run.font.name = "Arial"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(91, 100, 114)
        border = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "4")
        bottom.set(qn("w:color"), "D7DDE5")
        border.append(bottom)
        paragraph._p.get_or_add_pPr().append(border)

    def _configure_footer(self, section: Any) -> None:
        paragraph = section.footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = paragraph.add_run("Kimi AI")
        run.font.name = "Arial"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(91, 100, 114)
