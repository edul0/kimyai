from __future__ import annotations

import base64
import io
import re
import zipfile
from collections import Counter
from xml.etree import ElementTree as ET
from typing import Any

from docx import Document
from PIL import Image, ImageStat
from pypdf import PdfReader
from pptx import Presentation


def _compact_text(value: str, limit: int = 12000) -> str:
    text = " ".join(str(value or "").split())
    if len(text) <= limit:
        return text
    return text[: limit - 3].rstrip() + "..."


def _decode_data_url(content: str) -> tuple[str | None, bytes | None]:
    if not content.startswith("data:"):
        return None, None
    match = re.match(r"^data:([^;]+);base64,(.+)$", content, re.I | re.S)
    if not match:
        return None, None
    mime_type = match.group(1).strip().lower()
    try:
        payload = base64.b64decode(match.group(2), validate=False)
    except Exception:
        return mime_type, None
    return mime_type, payload


def _rgb_to_hex(color: tuple[int, int, int]) -> str:
    r, g, b = color
    return f"#{int(r):02X}{int(g):02X}{int(b):02X}"


def _brightness_label(value: float) -> str:
    if value < 75:
        return "muito escura"
    if value < 115:
        return "escura"
    if value < 165:
        return "equilibrada"
    if value < 210:
        return "clara"
    return "muito clara"


def _saturation_label(value: float) -> str:
    if value < 40:
        return "baixa"
    if value < 95:
        return "media"
    if value < 150:
        return "alta"
    return "muito alta"


def _extract_image_brief(data: bytes) -> tuple[str, dict[str, Any]]:
    try:
        image = Image.open(io.BytesIO(data)).convert("RGB")
    except Exception:
        return "", {}

    width, height = image.size
    if width > height:
        orientation = "paisagem"
    elif height > width:
        orientation = "retrato"
    else:
        orientation = "quadrada"

    sample = image.resize((min(220, width), min(220, height)))
    gray = sample.convert("L")
    hsv = sample.convert("HSV")
    gray_stat = ImageStat.Stat(gray)
    hsv_stat = ImageStat.Stat(hsv)
    brightness = float(gray_stat.mean[0]) if gray_stat.mean else 0.0
    contrast = float(gray_stat.stddev[0]) if gray_stat.stddev else 0.0
    saturation = float(hsv_stat.mean[1]) if hsv_stat.mean else 0.0

    quantized = sample.quantize(colors=5).convert("RGB")
    colors = quantized.getcolors(maxcolors=quantized.width * quantized.height) or []
    dominant = sorted(colors, key=lambda item: item[0], reverse=True)[:5]
    palette = [_rgb_to_hex(rgb) for _count, rgb in dominant]

    style_tags: list[str] = []
    if brightness < 100:
        style_tags.append("fundo escuro")
    if contrast >= 48:
        style_tags.append("contraste alto")
    if saturation >= 105:
        style_tags.append("cores vibrantes")
    elif saturation < 45:
        style_tags.append("paleta neutra")
    if orientation == "paisagem":
        style_tags.append("composicao horizontal")

    style_line = ", ".join(style_tags) if style_tags else "composicao limpa"
    summary = (
        f"Imagem {width}x{height} ({orientation}). "
        f"Paleta dominante: {', '.join(palette[:4]) or 'n/d'}. "
        f"Brilho medio: {brightness:.1f} ({_brightness_label(brightness)}). "
        f"Saturacao media: {saturation:.1f} ({_saturation_label(saturation)}). "
        f"Contraste estimado: {contrast:.1f}. "
        f"Sinais visuais: {style_line}."
    )
    metadata = {
        "width": width,
        "height": height,
        "orientation": orientation,
        "palette": palette,
        "brightness": round(brightness, 2),
        "saturation": round(saturation, 2),
        "contrast": round(contrast, 2),
        "style_tags": style_tags,
    }
    return _compact_text(summary, 1600), metadata


def _extract_pdf_text(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    parts: list[str] = []
    for page in reader.pages[:20]:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            continue
    return _compact_text("\n".join(parts), 18000)


def _extract_docx_text(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()]
    return _compact_text("\n".join(parts), 18000)


def _safe_xml_text(root: ET.Element | None, path: str, namespace: dict[str, str]) -> str:
    if root is None:
        return ""
    node = root.find(path, namespace)
    return (node.text or "").strip() if node is not None and node.text else ""


def _safe_xml_attr(root: ET.Element | None, path: str, attr: str, namespace: dict[str, str]) -> str:
    if root is None:
        return ""
    node = root.find(path, namespace)
    return (node.attrib.get(attr, "") or "").strip() if node is not None else ""


def _extract_pptx_theme(data: bytes) -> dict[str, str]:
    ns = {
        "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    }
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            theme_name = ""
            heading_font = ""
            body_font = ""
            accent_colors: list[str] = []

            if "ppt/theme/theme1.xml" in archive.namelist():
                root = ET.fromstring(archive.read("ppt/theme/theme1.xml"))
                theme_name = root.attrib.get("name", "").strip()
                heading_font = _safe_xml_attr(root, ".//a:themeElements/a:fontScheme/a:majorFont/a:latin", "typeface", ns)
                body_font = _safe_xml_attr(root, ".//a:themeElements/a:fontScheme/a:minorFont/a:latin", "typeface", ns)
                for tag in ("accent1", "accent2", "accent3", "accent4", "accent5", "accent6"):
                    node = root.find(f".//a:themeElements/a:clrScheme/a:{tag}", ns)
                    if node is None:
                        continue
                    srgb = node.find("./a:srgbClr", ns)
                    if srgb is not None:
                        value = srgb.attrib.get("val", "").strip()
                        if value:
                            accent_colors.append(f"#{value}")
            return {
                "theme_name": theme_name,
                "heading_font": heading_font,
                "body_font": body_font,
                "accent_colors": ", ".join(accent_colors[:4]),
            }
    except Exception:
        return {"theme_name": "", "heading_font": "", "body_font": "", "accent_colors": ""}


def _extract_pptx_thumbnail(data: bytes) -> tuple[str | None, bytes | None]:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names = archive.namelist()
            for candidate in ("docProps/thumbnail.jpeg", "docProps/thumbnail.jpg", "docProps/thumbnail.png"):
                if candidate in names:
                    mime = "image/jpeg" if candidate.endswith((".jpeg", ".jpg")) else "image/png"
                    return mime, archive.read(candidate)
    except Exception:
        pass
    return None, None


def _extract_pptx_text(data: bytes) -> tuple[str, dict[str, Any]]:
    prs = Presentation(io.BytesIO(data))
    layout_counter: Counter[str] = Counter()
    slide_sections: list[str] = []
    image_count = 0
    table_count = 0

    for index, slide in enumerate(list(prs.slides)[:20], start=1):
        title_text = ""
        try:
            if slide.shapes.title and slide.shapes.title.text:
                title_text = " ".join(slide.shapes.title.text.split())
        except Exception:
            title_text = ""

        text_parts: list[str] = []
        slide_image_count = 0
        slide_table_count = 0
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = " ".join((shape.text or "").split())
                if text:
                    text_parts.append(text)
            if getattr(shape, "shape_type", None) == 13:
                slide_image_count += 1
            if getattr(shape, "has_table", False):
                slide_table_count += 1

        image_count += slide_image_count
        table_count += slide_table_count
        layout_name = ""
        try:
            layout_name = (slide.slide_layout.name or "").strip()
        except Exception:
            layout_name = ""
        if layout_name:
            layout_counter[layout_name] += 1

        unique_text: list[str] = []
        seen: set[str] = set()
        for chunk in text_parts:
            if chunk not in seen:
                unique_text.append(chunk)
                seen.add(chunk)
        body_excerpt = " | ".join(unique_text[:4])
        label = title_text or f"Slide {index}"
        section = f"Slide {index}: {label}"
        if layout_name:
            section += f" [layout: {layout_name}]"
        if body_excerpt:
            section += f"\n{body_excerpt}"
        if slide_image_count or slide_table_count:
            markers: list[str] = []
            if slide_image_count:
                markers.append(f"{slide_image_count} imagem(ns)")
            if slide_table_count:
                markers.append(f"{slide_table_count} tabela(s)")
            section += f"\nElementos: {', '.join(markers)}"
        slide_sections.append(section)

    width = prs.slide_width
    height = prs.slide_height
    ratio = round(width / height, 3) if width and height else 0
    if 1.7 <= ratio <= 1.8:
        aspect = "16:9"
    elif 1.2 <= ratio <= 1.4:
        aspect = "4:3"
    else:
        aspect = f"{ratio}:1" if ratio else "desconhecido"

    theme = _extract_pptx_theme(data)
    layouts = ", ".join(f"{name} ({count})" for name, count in layout_counter.most_common(4))
    header_parts = [
        f"Apresentacao com {len(prs.slides)} slide(s).",
        f"Formato: {aspect}.",
    ]
    if layouts:
        header_parts.append(f"Layouts mais usados: {layouts}.")
    if theme.get("theme_name"):
        header_parts.append(f"Tema: {theme['theme_name']}.")
    if theme.get("heading_font") or theme.get("body_font"):
        header_parts.append(
            "Fontes do tema: "
            f"titulo={theme.get('heading_font') or 'n/d'}, corpo={theme.get('body_font') or 'n/d'}."
        )
    if theme.get("accent_colors"):
        header_parts.append(f"Cores de destaque: {theme['accent_colors']}.")
    if image_count or table_count:
        header_parts.append(f"Total de imagens: {image_count}. Total de tabelas: {table_count}.")

    text = "\n".join([" ".join(header_parts), *slide_sections])
    metadata = {
        "slides": len(prs.slides),
        "aspect_ratio": aspect,
        "layouts": dict(layout_counter),
        "images": image_count,
        "tables": table_count,
        **theme,
    }
    return _compact_text(text, 18000), metadata


def prepare_attachments(attachments: list[dict[str, Any]] | None) -> dict[str, Any]:
    prepared: list[dict[str, Any]] = []
    prompt_sections: list[str] = []
    visual_items: list[dict[str, Any]] = []

    for item in attachments or []:
        name = str(item.get("name") or "anexo").strip()[:180]
        declared_mime = str(item.get("mime_type") or "text/plain").strip().lower()
        kind = str(item.get("kind") or "text").strip().lower()
        content = str(item.get("content") or "")
        pptx_meta: dict[str, Any] = {}
        image_meta: dict[str, Any] = {}

        resolved_mime, binary = _decode_data_url(content)
        mime_type = resolved_mime or declared_mime
        extracted_text = ""
        notes: list[str] = []

        if mime_type == "application/pdf" and binary:
            extracted_text = _extract_pdf_text(binary)
            notes.append("pdf")
        elif mime_type == "application/vnd.openxmlformats-officedocument.presentationml.presentation" and binary:
            extracted_text, pptx_meta = _extract_pptx_text(binary)
            notes.extend(["pptx", "presentation"])
        elif mime_type in {
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "application/msword",
        } and binary:
            extracted_text = _extract_docx_text(binary)
            notes.append("docx")
        elif kind == "text":
            extracted_text = _compact_text(content, 18000)
            notes.append("text")
        elif mime_type.startswith("image/") and binary:
            extracted_text, image_meta = _extract_image_brief(binary)
            notes.append("image")
            if extracted_text:
                notes.append("image-brief")
            visual_items.append(
                {
                    "name": name,
                    "mime_type": mime_type,
                    "data": binary,
                    "kind": "image",
                }
            )
        elif binary:
            notes.append("binary")

        if mime_type == "application/vnd.openxmlformats-officedocument.presentationml.presentation" and binary:
            thumb_mime, thumb_bytes = _extract_pptx_thumbnail(binary)
            if thumb_bytes and thumb_mime:
                visual_items.append(
                    {
                        "name": f"{name}::thumbnail",
                        "mime_type": thumb_mime,
                        "data": thumb_bytes,
                        "kind": "image",
                    }
                )

        if mime_type == "application/pdf" and binary:
            visual_items.append(
                {
                    "name": name,
                    "mime_type": mime_type,
                    "data": binary,
                    "kind": "document",
                }
            )

        prepared_item = {
            "name": name,
            "mime_type": mime_type,
            "kind": kind,
            "bytes": binary,
            "text": extracted_text,
            "notes": notes,
            "metadata": {**pptx_meta, **image_meta},
            "visual": any(visual.get("name", "").startswith(name) for visual in visual_items),
        }
        prepared.append(prepared_item)

        if mime_type == "application/vnd.openxmlformats-officedocument.presentationml.presentation":
            prompt_sections.append(
                f"## {name} ({mime_type})\n"
                "PPTX anexado. Trate este arquivo como referencia de criacao: reaproveite narrativa, organizacao dos slides, tema, layouts, proporcao e estilo quando o usuario pedir continuidade, ajuste ou nova apresentacao baseada nele.\n"
                f"Resumo estrutural da apresentacao:\n{_compact_text(extracted_text, 7000)}"
            )
        elif mime_type.startswith("image/") and extracted_text:
            prompt_sections.append(
                f"## {name} ({mime_type})\n"
                "Imagem anexada. Use esta referencia visual para orientar estilo, composicao, paleta, hierarquia e atmosfera da resposta.\n"
                f"Resumo visual local do anexo:\n{_compact_text(extracted_text, 3000)}"
            )
        elif extracted_text:
            prompt_sections.append(
                f"## {name} ({mime_type})\nTexto extraido do anexo:\n{_compact_text(extracted_text, 6000)}"
            )
        elif mime_type.startswith("image/"):
            prompt_sections.append(
                f"## {name} ({mime_type})\nImagem anexada. Analise visual necessaria para descrever a cena e transcrever o texto visivel."
            )
        elif mime_type == "application/pdf":
            prompt_sections.append(
                f"## {name} ({mime_type})\nPDF anexado. Se a extracao local vier incompleta, complemente pela leitura visual do documento."
            )
        else:
            prompt_sections.append(f"## {name} ({mime_type})\nArquivo anexado.")

    return {
        "items": prepared,
        "prompt_context": "\n\n".join(prompt_sections).strip(),
        "visual_items": visual_items,
        "has_visual": bool(visual_items),
    }
