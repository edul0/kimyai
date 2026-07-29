"""Geradores determinísticos de artefatos profissionais da Kemy."""
from __future__ import annotations

import re
import zipfile
from pathlib import Path


def _hex(value: str, default: str) -> str:
    value = str(value or "").strip().lstrip("#")
    return value if re.fullmatch(r"[0-9a-fA-F]{6}", value) else default


def build_pptx(spec: dict, dest: Path) -> bool:
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt
    except Exception:
        return False
    primary = _hex(spec.get("primary"), "7657E8")
    accent = _hex(spec.get("accent"), "FF8B7B")
    dark = RGBColor(18, 17, 24)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slides = spec.get("slides") or []
    for index, item in enumerate(slides[:20]):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        bg = slide.background.fill
        bg.solid(); bg.fore_color.rgb = dark
        band = slide.shapes.add_shape(1, 0, 0, Inches(.18), prs.slide_height)
        band.fill.solid(); band.fill.fore_color.rgb = RGBColor.from_string(primary); band.line.fill.background()
        number = slide.shapes.add_textbox(Inches(11.8), Inches(.35), Inches(.8), Inches(.35))
        p = number.text_frame.paragraphs[0]; p.text = f"{index+1:02d}"; p.alignment = PP_ALIGN.RIGHT
        p.font.size = Pt(11); p.font.bold = True; p.font.color.rgb = RGBColor.from_string(accent)
        title = slide.shapes.add_textbox(Inches(.8), Inches(.7), Inches(11.2), Inches(1.25))
        p = title.text_frame.paragraphs[0]; p.text = str(item.get("title") or ""); p.font.size = Pt(30 if index else 38)
        p.font.bold = True; p.font.color.rgb = RGBColor(245, 243, 250)
        subtitle = str(item.get("subtitle") or "")
        if subtitle:
            box = slide.shapes.add_textbox(Inches(.82), Inches(1.75), Inches(10.7), Inches(.65))
            p = box.text_frame.paragraphs[0]; p.text = subtitle; p.font.size = Pt(15); p.font.color.rgb = RGBColor(175, 169, 190)
        body = slide.shapes.add_textbox(Inches(.9), Inches(2.55), Inches(11.3), Inches(3.9))
        tf = body.text_frame; tf.clear(); tf.word_wrap = True
        bullets = item.get("bullets") or []
        # Alterna lista editorial e cards para evitar uma apresentação inteira de bullets.
        if index > 0 and index % 3 == 0 and bullets:
            body.text = ""
            cols = min(3, len(bullets))
            for i, bullet in enumerate(bullets[:cols]):
                card = slide.shapes.add_shape(5, Inches(.85 + i * 4.08), Inches(2.65), Inches(3.65), Inches(2.45))
                card.fill.solid(); card.fill.fore_color.rgb = RGBColor(31, 28, 41)
                card.line.color.rgb = RGBColor.from_string(primary)
                card.text_frame.clear()
                p = card.text_frame.paragraphs[0]; p.text = f"{i+1:02d}"; p.font.size = Pt(13)
                p.font.bold = True; p.font.color.rgb = RGBColor.from_string(accent)
                p = card.text_frame.add_paragraph(); p.text = str(bullet); p.font.size = Pt(17)
                p.font.color.rgb = RGBColor(236, 232, 242); p.space_before = Pt(12)
        else:
            for i, bullet in enumerate(bullets[:6]):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = str(bullet); p.font.size = Pt(20); p.font.color.rgb = RGBColor(224, 220, 232)
                p.space_after = Pt(14); p.level = 0
        footer = slide.shapes.add_textbox(Inches(.82), Inches(6.95), Inches(7), Inches(.25))
        p = footer.text_frame.paragraphs[0]; p.text = str(spec.get("brand") or "Kemy")
        p.font.size = Pt(9); p.font.color.rgb = RGBColor(110, 105, 120)
    dest.parent.mkdir(parents=True, exist_ok=True)
    prs.save(dest)
    return bool(slides)


def build_docx(spec: dict, dest: Path) -> bool:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt, RGBColor
    except Exception:
        return False
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Inches(.7)
    sec.left_margin = sec.right_margin = Inches(.85)
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"; styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.name = "Aptos Display"; styles["Title"].font.size = Pt(32)
    styles["Title"].font.color.rgb = RGBColor.from_string(_hex(spec.get("primary"), "7657E8"))
    title = doc.add_heading(str(spec.get("title") or "Documento"), 0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if spec.get("subtitle"):
        p = doc.add_paragraph(str(spec["subtitle"])); p.style = styles["Subtitle"]
    for section in (spec.get("sections") or [])[:24]:
        doc.add_heading(str(section.get("title") or ""), level=1)
        for paragraph in (section.get("paragraphs") or [])[:8]:
            doc.add_paragraph(str(paragraph))
        for bullet in (section.get("bullets") or [])[:10]:
            doc.add_paragraph(str(bullet), style="List Bullet")
    footer = sec.footer.paragraphs[0]
    footer.text = str(spec.get("brand") or "Gerado com Kemy"); footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    dest.parent.mkdir(parents=True, exist_ok=True)
    doc.save(dest)
    return True


def build_logo_svg(spec: dict, dest: Path) -> bool:
    name = re.sub(r"[^\wÀ-ÿ -]", "", str(spec.get("brand") or "Kemy")).strip()[:28] or "Kemy"
    initials = "".join(part[0].upper() for part in name.split()[:2]) or "K"
    primary, accent = _hex(spec.get("primary"), "7657E8"), _hex(spec.get("accent"), "FF8B7B")
    svg = f'''<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1200 400" role="img" aria-label="{name}">
<defs><linearGradient id="g" x1="0" y1="0" x2="1" y2="1"><stop stop-color="#{primary}"/><stop offset="1" stop-color="#{accent}"/></linearGradient></defs>
<rect width="1200" height="400" rx="64" fill="#111018"/>
<rect x="72" y="72" width="256" height="256" rx="72" fill="url(#g)"/>
<text x="200" y="238" text-anchor="middle" font-family="Arial,sans-serif" font-size="112" font-weight="800" fill="white">{initials}</text>
<text x="382" y="220" font-family="Arial,sans-serif" font-size="96" font-weight="750" fill="#f6f3fa">{name}</text>
<path d="M386 258h520" stroke="url(#g)" stroke-width="12" stroke-linecap="round"/>
</svg>'''
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_text(svg, encoding="utf-8")
    return True


def build_pdf(spec: dict, dest: Path) -> bool:
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_LEFT
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer
    except Exception:
        return False
    primary = colors.HexColor("#" + _hex(spec.get("primary"), "7657E8"))
    styles = getSampleStyleSheet()
    title = ParagraphStyle("KemyTitle", parent=styles["Title"], fontName="Helvetica-Bold",
                           fontSize=28, leading=32, textColor=primary, alignment=TA_LEFT, spaceAfter=12)
    h1 = ParagraphStyle("KemyH1", parent=styles["Heading1"], fontName="Helvetica-Bold",
                        fontSize=17, leading=21, textColor=colors.HexColor("#211D2B"), spaceBefore=12, spaceAfter=8)
    body = ParagraphStyle("KemyBody", parent=styles["BodyText"], fontName="Helvetica",
                          fontSize=10.5, leading=15, textColor=colors.HexColor("#38333F"), spaceAfter=7)
    bullet = ParagraphStyle("KemyBullet", parent=body, leftIndent=12, firstLineIndent=-7, bulletIndent=0)
    dest.parent.mkdir(parents=True, exist_ok=True)

    def footer(canvas, doc):
        canvas.saveState(); canvas.setStrokeColor(primary); canvas.setLineWidth(1)
        canvas.line(18 * mm, 15 * mm, 192 * mm, 15 * mm)
        canvas.setFont("Helvetica", 8); canvas.setFillColor(colors.HexColor("#77717E"))
        canvas.drawString(18 * mm, 9 * mm, str(spec.get("brand") or "Gerado com Kemy"))
        canvas.drawRightString(192 * mm, 9 * mm, str(doc.page)); canvas.restoreState()

    story = [Paragraph(str(spec.get("title") or "Documento"), title)]
    if spec.get("subtitle"):
        story += [Paragraph(str(spec["subtitle"]), body), Spacer(1, 8 * mm)]
    for i, section in enumerate((spec.get("sections") or [])[:30]):
        if i and i % 5 == 0:
            story.append(PageBreak())
        story.append(Paragraph(str(section.get("title") or ""), h1))
        for paragraph in (section.get("paragraphs") or [])[:10]:
            story.append(Paragraph(str(paragraph), body))
        for item in (section.get("bullets") or [])[:12]:
            story.append(Paragraph("• " + str(item), bullet))
    doc = SimpleDocTemplate(str(dest), pagesize=A4, rightMargin=18*mm, leftMargin=18*mm,
                            topMargin=18*mm, bottomMargin=20*mm, title=str(spec.get("title") or "Documento"))
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return dest.exists() and dest.stat().st_size > 1000


def build_logo_kit(spec: dict, folder: Path) -> list[Path]:
    folder.mkdir(parents=True, exist_ok=True)
    outputs = []
    horizontal = folder / "logo-horizontal.svg"
    if build_logo_svg(spec, horizontal):
        outputs.append(horizontal)
    name = re.sub(r"[^\wÀ-ÿ -]", "", str(spec.get("brand") or "Kemy")).strip()[:28] or "Kemy"
    initials = "".join(part[0].upper() for part in name.split()[:2]) or "K"
    primary, accent = _hex(spec.get("primary"), "7657E8"), _hex(spec.get("accent"), "FF8B7B")
    icon = folder / "logo-icone.svg"
    icon.write_text(f'''<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 512 512"><defs><linearGradient id="g"><stop stop-color="#{primary}"/><stop offset="1" stop-color="#{accent}"/></linearGradient></defs><rect width="512" height="512" rx="144" fill="url(#g)"/><text x="256" y="320" text-anchor="middle" font-family="Arial" font-size="220" font-weight="800" fill="white">{initials}</text></svg>''', encoding="utf-8")
    mono = folder / "logo-monocromatico.svg"
    mono.write_text(f'''<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1000 260"><rect width="1000" height="260" rx="40" fill="white"/><text x="80" y="174" font-family="Arial" font-size="132" font-weight="800" fill="#111111">{name}</text></svg>''', encoding="utf-8")
    outputs += [icon, mono]
    guide = folder / "guia-da-marca.txt"
    guide.write_text(f"MARCA: {name}\nCOR PRIMÁRIA: #{primary}\nCOR DE DESTAQUE: #{accent}\nUSO: preserve área livre equivalente à altura das letras.\n", encoding="utf-8")
    outputs.append(guide)
    archive = folder / "kit-logo.zip"
    with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as zf:
        for path in outputs:
            zf.write(path, path.name)
    outputs.append(archive)
    return outputs


def validate_artifact(path: Path) -> dict:
    result = {"ok": False, "path": str(path), "size": 0, "detail": ""}
    try:
        result["size"] = path.stat().st_size
        if path.suffix.lower() in {".pptx", ".docx"}:
            with zipfile.ZipFile(path) as archive:
                names = archive.namelist()
                result["ok"] = "[Content_Types].xml" in names and len(names) > 5
                if path.suffix.lower() == ".pptx":
                    slide_count = sum(1 for n in names if re.match(r"ppt/slides/slide\d+\.xml$", n))
                    result["detail"] = f"{slide_count} slides"
        elif path.suffix.lower() == ".pdf":
            result["ok"] = path.read_bytes()[:4] == b"%PDF" and result["size"] > 1000
        elif path.suffix.lower() == ".svg":
            result["ok"] = "<svg" in path.read_text(encoding="utf-8", errors="ignore")[:500]
        else:
            result["ok"] = result["size"] > 0
    except Exception as exc:
        result["detail"] = str(exc)
    return result
